from __future__ import annotations

import hashlib
import shutil
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import numpy as np

from .baseline import DispatchSolution
from .parameters import Q1Parameters
from .time_index import build_official_q1_intervals


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def round4(value: float) -> float:
    return float(Decimal(str(float(value))).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))


def natural_day_four_hour_summaries(values: np.ndarray) -> tuple[float, ...]:
    if values.shape != (144,):
        raise ValueError("four-hour summary requires 144 template-ordered values")
    groups = (
        np.concatenate((values[143:144], values[0:23])),
        values[23:47],
        values[47:71],
        values[71:95],
        values[95:119],
        values[119:143],
    )
    if any(group.size != 24 for group in groups):
        raise AssertionError("natural-day aggregation did not produce six 24-slot groups")
    return tuple(float(np.sum(group)) for group in groups)


def export_result1_candidate(
    official_template: str | Path,
    candidate_path: str | Path,
    solution: DispatchSolution,
    params: Q1Parameters,
) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("openpyxl is required for result1 export") from exc

    source = Path(official_template).resolve()
    candidate = Path(candidate_path).resolve()
    if source == candidate:
        raise ValueError("candidate path must differ from the official template")
    if not source.is_file():
        raise FileNotFoundError(source)
    if candidate.exists():
        raise FileExistsError(f"refusing to overwrite existing candidate: {candidate}")
    if solution.G.shape != (params.interval_count,):
        raise ValueError("G does not match the official 144-slot template")
    original_hash_before = sha256_file(source)
    candidate.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, candidate)

    workbook = load_workbook(candidate, data_only=False)
    if workbook.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError(f"unexpected result1 sheets: {workbook.sheetnames}")
    plan = workbook["计划购电量"]
    expected = build_official_q1_intervals(params.interval_count, 10)
    for slot, interval in enumerate(expected, start=1):
        row = slot + 1
        if plan.cell(row, 1).value != interval.official_template_label:
            raise ValueError(f"official template label mismatch at slot {slot}")
        plan.cell(row, 2).value = round4(solution.G[slot - 1])
        plan.cell(row, 2).number_format = "0.0000"

    storage = workbook["充放电量"]
    expected_summary_labels = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]
    charge = natural_day_four_hour_summaries(solution.C)
    discharge = natural_day_four_hour_summaries(solution.D)
    for offset, label in enumerate(expected_summary_labels, start=2):
        if storage.cell(offset, 1).value != label:
            raise ValueError(f"storage summary label mismatch at row {offset}")
        storage.cell(offset, 2).value = round4(charge[offset - 2])
        storage.cell(offset, 3).value = round4(discharge[offset - 2])
        storage.cell(offset, 2).number_format = "0.0000"
        storage.cell(offset, 3).number_format = "0.0000"
    storage["E2"] = round4(solution.S[params.interval_count - 1])
    storage["E3"] = round4(solution.S[params.interval_count - 1])
    storage["E2"].number_format = "0.0000"
    storage["E3"].number_format = "0.0000"
    workbook.save(candidate)
    workbook.close()

    original_hash_after = sha256_file(source)
    if original_hash_after != original_hash_before:
        raise RuntimeError("official result1 template changed during export")
    return {
        "candidate_path": str(candidate),
        "candidate_sha256": sha256_file(candidate),
        "official_template_path": str(source),
        "official_sha256_before": original_hash_before,
        "official_sha256_after": original_hash_after,
        "official_template_unchanged": True,
    }

