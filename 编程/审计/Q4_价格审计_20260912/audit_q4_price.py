from __future__ import annotations

import argparse
import bisect
import hashlib
import json
import math
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl
import pandas as pd


DELTA = pd.Timedelta(minutes=10)
EXPECTED_MINUTES = tuple(range(10, 1441, 10))
FORMAL_START = date(2025, 2, 1)
FORMAL_END = date(2025, 12, 31)
VISIBILITY_RULES = (
    "TIMESTAMP_LE_DECISION_VISIBLE",
    "INTERVAL_END_LE_DECISION_VISIBLE",
)
MODE_ISSUES = {
    "Q4_2_DAILY_00": (0,),
    "Q4_3_ROLLING_4ISSUE_DIAGNOSTIC": (0, 360, 720, 1080),
}
PREDICTORS = ("P0_RECENT_SAME_CLOCK", "P1_WEEKDAY_SAME_CLOCK", "P2_EWMA", "P3_SHAPE_LEVEL")
EWMA_ALPHA = 0.35
EWMA_MAX_DAYS = 28  # retained for the guarded reference implementation below
P3_LEVEL_WINDOW_INTERVALS = 6


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def jsonable(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [jsonable(v) for v in value]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if not np.isfinite(value) else float(value)
    if isinstance(value, (pd.Timestamp, datetime, date, time)):
        return value.isoformat()
    if pd.isna(value):
        return None
    return value


def parse_header_minute(value: Any) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, datetime):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    if text == "0:00+1":
        return 1440
    hour_text, minute_text = text.split(":", 1)
    return int(hour_text) * 60 + int(minute_text)


def canonical_line(row: dict[str, Any]) -> str:
    return "|".join(
        (
            str(row["source_date"]),
            str(int(row["interval_index"])),
            str(int(row["sample_minute"])),
            pd.Timestamp(row["interval_start"]).isoformat(),
            pd.Timestamp(row["interval_end"]).isoformat(),
            f"{float(row['price_yuan_per_kwh']):.4f}",
        )
    )


def canonical_hash(rows: Iterable[dict[str, Any]]) -> str:
    digest = hashlib.sha256()
    for row in rows:
        digest.update(canonical_line(row).encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest().upper()


def load_raw(path: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet_names = list(wb.sheetnames)
    if sheet_names != ["Sheet1"]:
        raise AssertionError(f"unexpected Attachment4 sheets: {sheet_names}")
    ws = wb["Sheet1"]
    dimensions = [ws.max_row, ws.max_column]
    value_rows = ws.iter_rows(values_only=True)
    header = next(value_rows)
    header_minutes = tuple(parse_header_minute(value) for value in header[1:145])
    if len(header_minutes) != 144:
        raise AssertionError(f"unexpected Attachment4 time-column count: {len(header_minutes)}")
    formula_count = 0
    rows: list[dict[str, Any]] = []
    source_dates: list[date] = []
    for row_index, values in enumerate(value_rows, 2):
        raw_date = values[0]
        source_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        if not isinstance(source_date, date):
            raise AssertionError(f"invalid source date at row {row_index}: {raw_date!r}")
        source_dates.append(source_date)
        if len(values) < 145:
            raise AssertionError(f"short Attachment4 row at Excel row {row_index}: {len(values)} cells")
        for slot, (sample_minute, raw_value) in enumerate(zip(header_minutes, values[1:145]), 1):
            if isinstance(raw_value, str) and raw_value.startswith("="):
                formula_count += 1
            try:
                price = float(raw_value)
            except (TypeError, ValueError) as exc:
                raise AssertionError(f"invalid price at row={row_index}, slot={slot}: {raw_value!r}") from exc
            start = pd.Timestamp(source_date) + pd.Timedelta(minutes=sample_minute)
            rows.append(
                {
                    "source_date": source_date.isoformat(),
                    "interval_index": slot,
                    "sample_minute": sample_minute,
                    "interval_start": start,
                    "interval_end": start + DELTA,
                    "price_yuan_per_kwh": price,
                }
            )
    frame = pd.DataFrame(rows)
    info = {
        "sheet_names": sheet_names,
        "dimensions": dimensions,
        "header_first": str(ws.cell(1, 2).value),
        "header_last": str(ws.cell(1, 145).value),
        "header_minutes_match_expected": header_minutes == EXPECTED_MINUTES,
        "source_date_count": len(source_dates),
        "source_date_unique_count": len(set(source_dates)),
        "source_date_first": min(source_dates).isoformat(),
        "source_date_last": max(source_dates).isoformat(),
        "formula_count": formula_count,
        "row_count": len(frame),
    }
    wb.close()
    return frame, info


def load_normalized(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype={"source_date": str, "interval_start_date": str, "interval_end_date": str})
    required = {
        "source_date", "interval_index", "sample_minute", "interval_start_date",
        "interval_start_minute", "interval_end_date", "interval_end_minute", "price_yuan_per_kwh",
    }
    missing = required - set(frame.columns)
    if missing:
        raise AssertionError(f"normalized Attachment4 is missing columns: {sorted(missing)}")
    frame["interval_start"] = pd.to_datetime(frame["interval_start_date"]) + pd.to_timedelta(
        frame["interval_start_minute"], unit="m"
    )
    frame["interval_end"] = pd.to_datetime(frame["interval_end_date"]) + pd.to_timedelta(
        frame["interval_end_minute"], unit="m"
    )
    frame["interval_index"] = frame["interval_index"].astype(int)
    frame["sample_minute"] = frame["sample_minute"].astype(int)
    frame["price_yuan_per_kwh"] = pd.to_numeric(frame["price_yuan_per_kwh"], errors="coerce")
    return frame.sort_values(["source_date", "interval_index"]).reset_index(drop=True)


def reconcile(raw: pd.DataFrame, normalized: pd.DataFrame) -> dict[str, Any]:
    raw_sorted = raw.sort_values(["source_date", "interval_index"]).reset_index(drop=True)
    norm_sorted = normalized.sort_values(["source_date", "interval_index"]).reset_index(drop=True)
    keys = ["source_date", "interval_index"]
    merged = raw_sorted.merge(norm_sorted, on=keys, how="outer", suffixes=("_raw", "_norm"), indicator=True)
    both = merged["_merge"] == "both"
    checks = {
        "row_presence": both,
        "sample_minute": merged["sample_minute_raw"].eq(merged["sample_minute_norm"]),
        "interval_start": merged["interval_start_raw"].eq(merged["interval_start_norm"]),
        "interval_end": merged["interval_end_raw"].eq(merged["interval_end_norm"]),
        "price": np.isclose(
            merged["price_yuan_per_kwh_raw"], merged["price_yuan_per_kwh_norm"], rtol=0.0, atol=5e-12,
            equal_nan=False,
        ),
    }
    mismatch_by_field = {name: int((~pd.Series(mask).fillna(False)).sum()) for name, mask in checks.items()}
    mismatch_any = np.zeros(len(merged), dtype=bool)
    for mask in checks.values():
        mismatch_any |= ~np.asarray(pd.Series(mask).fillna(False), dtype=bool)
    raw_records = raw_sorted.to_dict("records")
    norm_records = norm_sorted.to_dict("records")
    return {
        "raw_row_count": len(raw_sorted),
        "normalized_row_count": len(norm_sorted),
        "joined_row_count": len(merged),
        "mismatch_by_field": mismatch_by_field,
        "mismatch_row_count": int(mismatch_any.sum()),
        "canonical_raw_sha256": canonical_hash(raw_records),
        "canonical_normalized_sha256": canonical_hash(norm_records),
        "canonical_hash_match": canonical_hash(raw_records) == canonical_hash(norm_records),
    }


def distribution_and_structure(frame: pd.DataFrame) -> dict[str, Any]:
    physical = frame.sort_values("interval_start").reset_index(drop=True).copy()
    prices = physical["price_yuan_per_kwh"].to_numpy(dtype=float)
    finite = np.isfinite(prices)
    valid = prices[finite]
    quantile_probs = [0.01, 0.05, 0.25, 0.50, 0.75, 0.95, 0.99]
    quantiles = np.quantile(valid, quantile_probs)
    q01, q05, q25, median, q75, q95, q99 = quantiles
    iqr = q75 - q25
    lower_outer = q25 - 3.0 * iqr
    upper_outer = q75 + 3.0 * iqr
    differences = np.diff(valid)
    abs_diff = np.abs(differences)
    diff_median = float(np.median(abs_diff))
    diff_mad = float(np.median(np.abs(abs_diff - diff_median)))
    local_threshold = diff_median + 6.0 * diff_mad
    starts = physical["interval_start"]
    gaps = starts.diff().dropna()
    expected_count = int((starts.max() - starts.min()) / DELTA) + 1
    return {
        "count": int(len(prices)),
        "finite_count": int(finite.sum()),
        "nonfinite_count": int((~finite).sum()),
        "min": float(valid.min()),
        "max": float(valid.max()),
        "mean": float(valid.mean()),
        "median": float(median),
        "std_sample": float(valid.std(ddof=1)),
        "P01": float(q01),
        "P05": float(q05),
        "P25": float(q25),
        "P75": float(q75),
        "P95": float(q95),
        "P99": float(q99),
        "negative_price_count": int((valid < 0).sum()),
        "zero_price_count": int((valid == 0).sum()),
        "strictly_positive_price_count": int((valid > 0).sum()),
        "duplicate_source_slot_count": int(frame.duplicated(["source_date", "interval_index"]).sum()),
        "duplicate_interval_start_count": int(frame.duplicated(["interval_start"]).sum()),
        "physical_first_interval_start": starts.min().isoformat(),
        "physical_last_interval_start": starts.max().isoformat(),
        "physical_last_interval_end": physical["interval_end"].max().isoformat(),
        "expected_contiguous_row_count": expected_count,
        "missing_physical_interval_count_within_observed_span": expected_count - len(starts.unique()),
        "non_10_minute_gap_count": int((gaps != DELTA).sum()),
        "tukey_outer_lower_fence": float(lower_outer),
        "tukey_outer_upper_fence": float(upper_outer),
        "lower_outer_fence_count": int((valid < lower_outer).sum()),
        "upper_outer_fence_count": int((valid > upper_outer).sum()),
        "max_abs_adjacent_change": float(abs_diff.max()),
        "P99_abs_adjacent_change": float(np.quantile(abs_diff, 0.99)),
        "local_jump_threshold_median_plus_6MAD": float(local_threshold),
        "local_jump_count": int((abs_diff > local_threshold).sum()),
    }


def add_physical_features(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.sort_values("interval_start").reset_index(drop=True).copy()
    result["clock_minute"] = result["interval_start"].dt.hour * 60 + result["interval_start"].dt.minute
    result["clock_label"] = result["interval_start"].dt.strftime("%H:%M")
    result["physical_date"] = result["interval_start"].dt.date.astype(str)
    result["weekday"] = result["interval_start"].dt.weekday
    result["weekday_name"] = result["interval_start"].dt.day_name()
    result["month"] = result["interval_start"].dt.month
    return result


def profiles(frame: pd.DataFrame, output_dir: Path) -> dict[str, Any]:
    physical = add_physical_features(frame)
    price = "price_yuan_per_kwh"
    aggregations = {"count": (price, "size"), "mean": (price, "mean"), "median": (price, "median"), "std": (price, "std"), "min": (price, "min"), "max": (price, "max")}
    intraday = physical.groupby(["clock_minute", "clock_label"], as_index=False).agg(**aggregations)
    weekday = physical.groupby(["weekday", "weekday_name"], as_index=False).agg(**aggregations).sort_values("weekday")
    month = physical.groupby("month", as_index=False).agg(**aggregations).sort_values("month")
    intraday.to_csv(output_dir / "price_profile_intraday.csv", index=False, encoding="utf-8-sig")
    weekday.to_csv(output_dir / "price_profile_weekday.csv", index=False, encoding="utf-8-sig")
    month.to_csv(output_dir / "price_profile_month.csv", index=False, encoding="utf-8-sig")
    high = intraday.nlargest(5, "mean")[["clock_label", "mean"]].to_dict("records")
    low = intraday.nsmallest(5, "mean")[["clock_label", "mean"]].to_dict("records")
    return {
        "highest_mean_intraday_clocks": high,
        "lowest_mean_intraday_clocks": low,
        "weekday_mean_range": [float(weekday["mean"].min()), float(weekday["mean"].max())],
        "month_mean_range": [float(month["mean"].min()), float(month["mean"].max())],
        "highest_mean_weekday": weekday.loc[weekday["mean"].idxmax(), ["weekday_name", "mean"]].to_dict(),
        "lowest_mean_weekday": weekday.loc[weekday["mean"].idxmin(), ["weekday_name", "mean"]].to_dict(),
        "highest_mean_month": month.loc[month["mean"].idxmax(), ["month", "mean"]].to_dict(),
        "lowest_mean_month": month.loc[month["mean"].idxmin(), ["month", "mean"]].to_dict(),
    }


def stability(frame: pd.DataFrame, output_dir: Path) -> dict[str, Any]:
    physical = frame.sort_values("interval_start").reset_index(drop=True)
    series = physical.set_index("interval_start")["price_yuan_per_kwh"].astype(float)
    rows: list[dict[str, Any]] = []
    for label, lag in (("10min", 1), ("1hour", 6), ("1day_same_clock", 144), ("7day_same_weekday_clock", 1008)):
        current = series.iloc[lag:].to_numpy()
        previous = series.iloc[:-lag].to_numpy()
        delta = current - previous
        rows.append(
            {
                "lag_label": label,
                "lag_intervals": lag,
                "pair_count": len(delta),
                "pearson_correlation": float(np.corrcoef(current, previous)[0, 1]),
                "mae": float(np.mean(np.abs(delta))),
                "rmse": float(np.sqrt(np.mean(delta ** 2))),
                "bias_current_minus_previous": float(np.mean(delta)),
                "p90_absolute_change": float(np.quantile(np.abs(delta), 0.9)),
            }
        )
    pd.DataFrame(rows).to_csv(output_dir / "price_stability.csv", index=False, encoding="utf-8-sig")
    return {row["lag_label"]: row for row in rows}


@dataclass(frozen=True)
class HistRecord:
    start: pd.Timestamp
    price: float


class PriceHistory:
    def __init__(self, frame: pd.DataFrame) -> None:
        physical = add_physical_features(frame)
        self.lookup = {
            (row.interval_start.date(), int(row.clock_minute)): float(row.price_yuan_per_kwh)
            for row in physical.itertuples()
        }
        self.global_records = tuple(
            HistRecord(row.interval_start, float(row.price_yuan_per_kwh)) for row in physical.itertuples()
        )
        self.global_starts = tuple(row.start for row in self.global_records)

    @staticmethod
    def cutoff(decision: pd.Timestamp, rule: str) -> pd.Timestamp:
        if rule == "TIMESTAMP_LE_DECISION_VISIBLE":
            return decision
        if rule == "INTERVAL_END_LE_DECISION_VISIBLE":
            return decision - DELTA
        raise ValueError(rule)

    def latest_any(self, cutoff: pd.Timestamp) -> float | None:
        index = bisect.bisect_right(self.global_starts, cutoff) - 1
        return None if index < 0 else self.global_records[index].price

    @staticmethod
    def _clock(target: pd.Timestamp) -> int:
        return target.hour * 60 + target.minute

    def _candidate_date(self, clock: int, cutoff: pd.Timestamp) -> date:
        cutoff_clock = cutoff.hour * 60 + cutoff.minute
        return cutoff.date() if clock <= cutoff_clock else cutoff.date() - timedelta(days=1)

    def p0(self, target: pd.Timestamp, cutoff: pd.Timestamp) -> float | None:
        clock = self._clock(target)
        candidate = self._candidate_date(clock, cutoff)
        for _ in range(370):
            value = self.lookup.get((candidate, clock))
            if value is not None:
                return value
            candidate -= timedelta(days=1)
        return self.latest_any(cutoff)

    def p1(self, target: pd.Timestamp, cutoff: pd.Timestamp) -> float | None:
        clock = self._clock(target)
        target_weekday = target.weekday()
        candidate = self._candidate_date(clock, cutoff)
        for _ in range(370):
            if candidate.weekday() == target_weekday:
                value = self.lookup.get((candidate, clock))
                if value is not None:
                    return value
            candidate -= timedelta(days=1)
        return self.p0(target, cutoff)

    def ewma(self, target: pd.Timestamp, cutoff: pd.Timestamp) -> float | None:
        clock = self._clock(target)
        candidate = self._candidate_date(clock, cutoff)
        newest_first: list[float] = []
        for _ in range(370):
            value = self.lookup.get((candidate, clock))
            if value is not None:
                newest_first.append(value)
                if len(newest_first) == EWMA_MAX_DAYS:
                    break
            candidate -= timedelta(days=1)
        if not newest_first:
            return self.latest_any(cutoff)
        values = np.asarray(list(reversed(newest_first)), dtype=float)
        ages = np.arange(len(values) - 1, -1, -1, dtype=float)
        weights = (1.0 - EWMA_ALPHA) ** ages
        return float(np.dot(values, weights) / weights.sum())

    def ewma_before(self, target: pd.Timestamp, exclusive: pd.Timestamp) -> float | None:
        return self.ewma(target, exclusive - DELTA)

    def p3_level(self, decision: pd.Timestamp, cutoff: pd.Timestamp) -> float:
        day_start = decision.normalize()
        if cutoff < day_start:
            return 1.0
        cutoff_clock = cutoff.hour * 60 + cutoff.minute
        observed_today = [
            HistRecord(day_start + pd.Timedelta(minutes=clock), self.lookup[(day_start.date(), clock)])
            for clock in range(0, cutoff_clock + 1, 10)
            if (day_start.date(), clock) in self.lookup
        ]
        ratios: list[float] = []
        for record in observed_today[-P3_LEVEL_WINDOW_INTERVALS:]:
            shape = self.ewma_before(record.start, day_start)
            if shape is not None and shape > 0 and np.isfinite(shape):
                ratios.append(record.price / shape)
        return 1.0 if not ratios else float(np.median(ratios))


@dataclass
class MetricAccumulator:
    errors: list[float] = field(default_factory=list)
    abs_errors: list[float] = field(default_factory=list)
    rank_correlations: list[float] = field(default_factory=list)
    peak_recalls: list[float] = field(default_factory=list)
    offpeak_recalls: list[float] = field(default_factory=list)
    prediction_min: float = math.inf
    prediction_max: float = -math.inf
    negative_count: int = 0
    nonfinite_count: int = 0
    observed_at_decision_count: int = 0
    event_count: int = 0

    @staticmethod
    def _average_ranks(values: np.ndarray) -> np.ndarray:
        _, inverse, counts = np.unique(values, return_inverse=True, return_counts=True)
        ends = np.cumsum(counts)
        starts = ends - counts
        average_by_unique = 0.5 * (starts + ends - 1) + 1.0
        return average_by_unique[inverse]

    def add(
        self,
        actual: np.ndarray,
        prediction: np.ndarray,
        observed: np.ndarray,
        *,
        compute_event_metrics: bool,
    ) -> None:
        if len(actual) == 0:
            return
        finite = np.isfinite(prediction)
        self.nonfinite_count += int((~finite).sum())
        if not finite.all():
            return
        error = prediction - actual
        self.errors.extend(error.tolist())
        self.abs_errors.extend(np.abs(error).tolist())
        self.prediction_min = min(self.prediction_min, float(prediction.min()))
        self.prediction_max = max(self.prediction_max, float(prediction.max()))
        self.negative_count += int((prediction < 0).sum())
        self.observed_at_decision_count += int(observed.sum())
        self.event_count += 1
        if compute_event_metrics:
            actual_rank = self._average_ranks(actual)
            pred_rank = self._average_ranks(prediction)
            rank_corr = np.corrcoef(actual_rank, pred_rank)[0, 1]
            if np.isfinite(rank_corr):
                self.rank_correlations.append(float(rank_corr))
            k = max(1, int(math.ceil(0.10 * len(actual))))
            actual_top = set(np.argpartition(actual, -k)[-k:])
            pred_top = set(np.argpartition(prediction, -k)[-k:])
            actual_bottom = set(np.argpartition(actual, k - 1)[:k])
            pred_bottom = set(np.argpartition(prediction, k - 1)[:k])
            self.peak_recalls.append(len(actual_top & pred_top) / k)
            self.offpeak_recalls.append(len(actual_bottom & pred_bottom) / k)

    def finish(self) -> dict[str, Any]:
        errors = np.asarray(self.errors, dtype=float)
        absolute = np.asarray(self.abs_errors, dtype=float)
        return {
            "prediction_count": int(len(errors)),
            "decision_event_count": self.event_count,
            "MAE": float(absolute.mean()),
            "RMSE": float(np.sqrt(np.mean(errors ** 2))),
            "bias_prediction_minus_actual": float(errors.mean()),
            "P90_absolute_error": float(np.quantile(absolute, 0.90)),
            "mean_same_day_rank_correlation": None if not self.rank_correlations else float(np.mean(self.rank_correlations)),
            "median_same_day_rank_correlation": None if not self.rank_correlations else float(np.median(self.rank_correlations)),
            "mean_peak_10pct_recall": None if not self.peak_recalls else float(np.mean(self.peak_recalls)),
            "mean_offpeak_10pct_recall": None if not self.offpeak_recalls else float(np.mean(self.offpeak_recalls)),
            "prediction_min": self.prediction_min,
            "prediction_max": self.prediction_max,
            "negative_prediction_count": self.negative_count,
            "nonfinite_prediction_count": self.nonfinite_count,
            "observed_at_decision_count": self.observed_at_decision_count,
        }


def prediction_diagnostics(frame: pd.DataFrame, output_dir: Path) -> list[dict[str, Any]]:
    # Strictly lagged same-clock candidates are precomputed once.  The event
    # loop below only reveals a current timestamp when the selected visibility
    # rule legally allows it.
    physical = add_physical_features(frame).sort_values(["clock_minute", "interval_start"]).copy()
    grouped = physical.groupby("clock_minute", group_keys=False)
    physical["p0"] = grouped["price_yuan_per_kwh"].shift(1)
    physical["p0_shift2"] = grouped["price_yuan_per_kwh"].shift(2)
    physical["p1"] = grouped["price_yuan_per_kwh"].shift(7)
    physical["p2"] = grouped["price_yuan_per_kwh"].transform(
        lambda series: series.shift(1).ewm(alpha=EWMA_ALPHA, adjust=True, min_periods=1).mean()
    )
    physical["p2_before_previous"] = grouped["p2"].shift(1)
    physical["p1"] = physical["p1"].fillna(physical["p0"])
    predictor_columns = physical[
        ["interval_start", "p0", "p0_shift2", "p1", "p2", "p2_before_previous"]
    ].set_index("interval_start")
    physical_by_date = {
        day_value: group.sort_values("interval_start").reset_index(drop=True)
        for day_value, group in physical.groupby(physical["interval_start"].dt.date)
    }
    enriched = frame.merge(
        predictor_columns,
        left_on="interval_start",
        right_index=True,
        how="left",
        validate="one_to_one",
    )
    actual_by_source_date = {
        source_date: group.sort_values("interval_index").reset_index(drop=True)
        for source_date, group in enriched.groupby("source_date")
    }
    accumulators: dict[tuple[str, str, str, str], MetricAccumulator] = {}
    for rule in VISIBILITY_RULES:
        for source_day in pd.date_range(FORMAL_START, FORMAL_END, freq="D"):
            group = actual_by_source_date[source_day.date().isoformat()]
            for issue_minute in MODE_ISSUES["Q4_3_ROLLING_4ISSUE_DIAGNOSTIC"]:
                modes = ["Q4_3_ROLLING_4ISSUE_DIAGNOSTIC"]
                if issue_minute == 0:
                    modes.append("Q4_2_DAILY_00")
                decision = source_day + pd.Timedelta(minutes=issue_minute)
                cutoff = decision if rule == "TIMESTAMP_LE_DECISION_VISIBLE" else decision - DELTA
                remaining = group[group["interval_start"] >= decision]
                if remaining.empty:
                    continue
                targets = pd.to_datetime(remaining["interval_start"])
                actual = remaining["price_yuan_per_kwh"].to_numpy(dtype=float)
                observed = (targets <= cutoff).to_numpy(dtype=bool)
                p0 = remaining["p0"].to_numpy(dtype=float, copy=True)
                p1 = remaining["p1"].to_numpy(dtype=float, copy=True)
                p2 = remaining["p2"].to_numpy(dtype=float, copy=True)

                # At 00:00 under interval-end visibility, the just-started
                # 00:00 interval is not yet visible. Slot 144 targets the next
                # midnight and therefore needs one additional daily lag.
                if rule == "INTERVAL_END_LE_DECISION_VISIBLE" and issue_minute == 0:
                    next_midnight = (targets.dt.date > source_day.date()).to_numpy(dtype=bool)
                    p0[next_midnight] = remaining.loc[next_midnight, "p0_shift2"].to_numpy(dtype=float)
                    p2[next_midnight] = remaining.loc[next_midnight, "p2_before_previous"].to_numpy(dtype=float)

                physical_day = physical_by_date[source_day.date()]
                visible_today = physical_day[
                    physical_day["interval_start"] <= cutoff
                ].tail(P3_LEVEL_WINDOW_INTERVALS)
                ratios = (
                    visible_today["price_yuan_per_kwh"].to_numpy(dtype=float)
                    / visible_today["p2"].to_numpy(dtype=float)
                )
                ratios = ratios[np.isfinite(ratios) & (ratios > 0)]
                level = 1.0 if len(ratios) == 0 else float(np.median(ratios))
                predictions = {
                    "P0_RECENT_SAME_CLOCK": p0,
                    "P1_WEEKDAY_SAME_CLOCK": p1,
                    "P2_EWMA": p2,
                    "P3_SHAPE_LEVEL": p2 * level,
                }
                for prediction in predictions.values():
                    prediction[observed] = actual[observed]
                    if not np.isfinite(prediction).all():
                        raise AssertionError("formal-period predictor unexpectedly has no visible price history")

                for mode in modes:
                    for name, prediction in predictions.items():
                        for scope, scope_mask in (
                            ("ALL_PLANNING_PRICE_INPUTS", np.ones(len(actual), dtype=bool)),
                            ("UNKNOWN_AT_DECISION_ONLY", ~observed),
                        ):
                            key = (rule, mode, name, scope)
                            accumulators.setdefault(key, MetricAccumulator()).add(
                                actual[scope_mask],
                                prediction[scope_mask],
                                observed[scope_mask],
                                compute_event_metrics=scope == "ALL_PLANNING_PRICE_INPUTS",
                            )
    rows: list[dict[str, Any]] = []
    for (rule, mode, predictor, scope), accumulator in sorted(accumulators.items()):
        row = {
            "visibility_rule": rule,
            "diagnostic_mode": mode,
            "predictor": predictor,
            "evaluation_scope": scope,
            **accumulator.finish(),
        }
        rows.append(row)
    pd.DataFrame(rows).to_csv(output_dir / "prediction_metrics.csv", index=False, encoding="utf-8-sig")
    return rows


def boundary_rows(frame: pd.DataFrame, output_dir: Path) -> list[dict[str, Any]]:
    wanted = (
        ("2025-01-01", 1), ("2025-01-01", 143), ("2025-01-01", 144),
        ("2025-12-31", 1), ("2025-12-31", 143), ("2025-12-31", 144),
    )
    rows: list[dict[str, Any]] = []
    for source_date, slot in wanted:
        match = frame[(frame["source_date"] == source_date) & (frame["interval_index"] == slot)]
        if len(match) != 1:
            raise AssertionError(f"boundary key missing or duplicate: {source_date}, {slot}")
        row = match.iloc[0]
        rows.append(
            {
                "source_date": source_date,
                "template_slot": slot,
                "sample_minute": int(row["sample_minute"]),
                "interval_start": row["interval_start"].isoformat(),
                "interval_end": row["interval_end"].isoformat(),
                "price_yuan_per_kwh": float(row["price_yuan_per_kwh"]),
            }
        )
    pd.DataFrame(rows).to_csv(output_dir / "boundary_examples.csv", index=False, encoding="utf-8-sig")
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description="Independent read-only Attachment4 and causal price diagnostic audit")
    parser.add_argument("--raw-xlsx", type=Path, required=True)
    parser.add_argument("--normalized-csv", type=Path, required=True)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    print("[Q4 price audit] loading raw workbook", flush=True)
    raw, raw_info = load_raw(args.raw_xlsx)
    print("[Q4 price audit] loading normalized CSV", flush=True)
    normalized = load_normalized(args.normalized_csv)
    print("[Q4 price audit] reconciling raw and normalized data", flush=True)
    reconciliation = reconcile(raw, normalized)
    distribution = distribution_and_structure(normalized)
    print("[Q4 price audit] computing distribution profiles", flush=True)
    profile_summary = profiles(normalized, args.output_dir)
    stability_summary = stability(normalized, args.output_dir)
    boundaries = boundary_rows(normalized, args.output_dir)
    print("[Q4 price audit] evaluating causal price candidates", flush=True)
    prediction_rows = prediction_diagnostics(normalized, args.output_dir)
    print("[Q4 price audit] writing summary", flush=True)

    manifest_hash = None
    manifest_match = None
    if args.manifest:
        manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
        entry = next((item for item in manifest if item.get("relative_path") == "附件\\附件4.xlsx"), None)
        manifest_hash = None if entry is None else entry.get("sha256")
        manifest_match = manifest_hash == sha256(args.raw_xlsx)

    summary = {
        "audit_id": "Q4-PRICE-AUDIT-20260912",
        "status": "READ_ONLY_AUDIT_COMPLETE_VISIBILITY_RULE_UNRESOLVED",
        "solver_executed": False,
        "formal_q4_dispatch_executed": False,
        "inputs": {
            "raw_xlsx": str(args.raw_xlsx.resolve()),
            "raw_xlsx_sha256": sha256(args.raw_xlsx),
            "normalized_csv": str(args.normalized_csv.resolve()),
            "normalized_csv_sha256": sha256(args.normalized_csv),
            "manifest": None if args.manifest is None else str(args.manifest.resolve()),
            "manifest_attachment4_sha256": manifest_hash,
            "manifest_raw_hash_match": manifest_match,
        },
        "raw_structure": raw_info,
        "raw_normalized_reconciliation": reconciliation,
        "distribution_and_structure": distribution,
        "profiles": profile_summary,
        "stability": stability_summary,
        "boundary_examples": boundaries,
        "time_semantics": {
            "raw_labels": "0:10, 0:20, ..., 23:50, 0:00+1",
            "audited_mapping": "instantaneous timestamp is the left endpoint of the following 10-minute interval",
            "template_day_span": "[source_date 00:10, source_date+1 day 00:10)",
            "natural_day_mapping": "natural day d uses source_date d-1 slot144 for [d 00:00,d 00:10) plus source_date d slots1-143",
            "jan1_natural_0000_gap": "2025-01-01 00:00-00:10 is outside Attachment4 physical coverage; template-day Jan1 starts at 00:10",
            "last_coverage": "source_date 2025-12-31 slot144 covers 2026-01-01 00:00-00:10",
        },
        "visibility_audit": {
            "conclusion": "BLOCKED",
            "reason": "official statement and workbook identify real-time fluctuating prices and timestamps but do not state publication/availability at interval start versus after interval completion",
            "diagnosed_rules": list(VISIBILITY_RULES),
        },
        "prediction_preregistration": {
            "formal_diagnostic_period": [FORMAL_START.isoformat(), FORMAL_END.isoformat()],
            "warmup_period": ["2025-01-01", "2025-01-31"],
            "P0": "latest legally visible actual price at the same natural clock; fallback latest visible price",
            "P1": "latest legally visible same-weekday same-clock actual; fallback P0",
            "P2": f"strictly lagged same-clock expanding EWMA alpha={EWMA_ALPHA}; fallback latest visible price when required",
            "P3": f"P2 daily shape multiplied by median actual/shape ratio from the latest {P3_LEVEL_WINDOW_INTERVALS} legally visible same-day intervals; no additive clipping",
            "nonnegative_design": "all four candidates are positive combinations/selections of strictly positive observed prices; known-at-decision target prices are passed through as observed",
            "model_selection_warning": "prediction diagnostics do not select the optimization model; final ranking requires downstream realized settlement cost under a preregistered rolling experiment",
        },
        "prediction_metrics": prediction_rows,
    }
    output_path = args.output_dir / "summary.json"
    output_path.write_text(json.dumps(jsonable(summary), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(output_path.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
