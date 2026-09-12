from __future__ import annotations

import hashlib
import json
import platform
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import Q43Config


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def core_python_hashes() -> dict[str, str]:
    root = Path(__file__).resolve().parent
    return {path.name: sha256_file(path) for path in sorted(root.glob("*.py"))}


def template_audit(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheets = [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column} for ws in workbook.worksheets]
    formulas = sum(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for ws in workbook.worksheets for row in ws.iter_rows() for cell in row
    )
    workbook.close()
    expected = [
        {"name": "计划购电量", "rows": 335, "columns": 147},
        {"name": "调整购电量", "rows": 335, "columns": 147},
        {"name": "充放电量", "rows": 26, "columns": 6},
        {"name": "紧急购电量", "rows": 11, "columns": 3},
    ]
    return {"sheets": sheets, "formula_count": formulas, "matches_locked_layout": sheets == expected and formulas == 0}


def _git_identity() -> dict[str, Any]:
    repo = Path(__file__).resolve().parents[2]
    def call(*args: str) -> str:
        result = subprocess.run(
            ["git", "-c", f"safe.directory={repo}", "-C", str(repo), *args],
            capture_output=True, text=True, check=False,
        )
        return result.stdout.strip() if result.returncode == 0 else f"UNAVAILABLE:{result.stderr.strip()}"
    dirty = call("status", "--porcelain")
    return {"git_head": call("rev-parse", "HEAD"), "git_dirty": bool(dirty) if not dirty.startswith("UNAVAILABLE") else dirty}


def verify_inputs(config: Q43Config) -> dict[str, Any]:
    required = {
        "official_result4_3": config.official_result4_3_template,
        "attachment2": config.q3.normalized_actual_input,
        "attachment3_interp": config.q3.attachment3_interp_input,
        "attachment3_mapping_manifest": config.q3.attachment3_mapping_manifest,
        "attachment4": config.q4.normalized_price_input,
        "frozen_q3_manifest": config.frozen_q3_run / "run_manifest.json",
        "frozen_q3_dispatch": config.frozen_q3_run / "dispatch_timeseries.csv",
        "frozen_q4_2_manifest": config.frozen_q4_2_run / "run_manifest.json",
    }
    missing = [str(path) for path in required.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Q4-3 required inputs missing: {missing}")
    hashes = {key: sha256_file(path) for key, path in required.items()}
    locked = {
        "official_result4_3": config.official_result4_3_sha256,
        "attachment2": config.q4.normalized_actual_sha256,
        "attachment4": config.q4.normalized_price_sha256,
        "frozen_q3_manifest": config.frozen_q3_manifest_sha256,
        "frozen_q3_dispatch": config.frozen_q3_dispatch_sha256,
        "frozen_q4_2_manifest": config.frozen_q4_2_manifest_sha256,
    }
    differences = {key: {"expected": value, "actual": hashes[key]} for key, value in locked.items() if hashes[key] != value}
    if differences:
        raise RuntimeError(f"Q4_3_IDENTITY_HARD_FAIL:{differences}")
    q3_manifest = json.loads(required["frozen_q3_manifest"].read_text(encoding="utf-8"))
    q4_manifest = json.loads(required["frozen_q4_2_manifest"].read_text(encoding="utf-8"))
    if q3_manifest.get("track") != config.q3_frozen_track or q3_manifest.get("model_version") != config.q3_frozen_model_version:
        raise RuntimeError("Q4_3_FROZEN_Q3_IDENTITY_HARD_FAIL")
    if q4_manifest.get("model_version") != config.q4_2_frozen_model_version or q4_manifest.get("track") != "Q4_2_PRICE_CANDIDATE_P2":
        raise RuntimeError("Q4_3_FROZEN_Q4_2_IDENTITY_HARD_FAIL")
    template = template_audit(config.official_result4_3_template)
    if not template["matches_locked_layout"]:
        raise RuntimeError("Q4_3_TEMPLATE_HARD_FAIL")
    return {
        "status": "PASS", "input_sha256": hashes, "template_audit": template,
        "frozen_q2_identity": q3_manifest.get("integrity", {}).get("q2_reference_manifest_sha256"),
        "frozen_q3_identity": {"manifest_sha256": hashes["frozen_q3_manifest"], "dispatch_sha256": hashes["frozen_q3_dispatch"]},
        "q4_2_predictor_spec_identity": {"manifest_sha256": hashes["frozen_q4_2_manifest"], "predictor": config.q4_2_frozen_predictor},
    }


def identity_snapshot(config: Q43Config, config_path: Path, inputs: dict[str, Any], feb1_soc: float, feb1_source: dict[str, object], highs_version: str) -> dict[str, Any]:
    return {
        "q4_3_core_python_sha256": core_python_hashes(),
        "config_sha256": sha256_file(config_path.resolve()),
        "input_sha256": inputs["input_sha256"],
        "frozen_q2_identity": inputs["frozen_q2_identity"],
        "frozen_q3_identity": inputs["frozen_q3_identity"],
        "q4_2_predictor_spec_identity": inputs["q4_2_predictor_spec_identity"],
        "visibility_rule": config.visibility_rule,
        "pv_issue_set": [0, 360, 720, 1080],
        "pv_mapping_method": "INTERP",
        "cost_semantics": "MODEL_B",
        "formal_dates": [config.output_start.isoformat(), config.output_end.isoformat()],
        "jan_warmup_rule": config.warmup_rule,
        "feb1_initial_soc": feb1_soc,
        "feb1_initial_soc_identity": feb1_source,
        "python_version": sys.version,
        "highs_version": highs_version,
        "platform": platform.platform(),
        **_git_identity(),
    }


def assert_identity_unchanged(before: dict[str, Any], config: Q43Config, config_path: Path, inputs: dict[str, Any], feb1_soc: float, feb1_source: dict[str, object], highs_version: str) -> None:
    after = identity_snapshot(config, config_path, inputs, feb1_soc, feb1_source, highs_version)
    if after != before:
        raise RuntimeError("Q4_3_IDENTITY_HARD_FAIL: code/config/input/environment changed during run")

