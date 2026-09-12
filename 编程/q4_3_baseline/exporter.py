from __future__ import annotations

from dataclasses import replace
from datetime import date
from pathlib import Path
from typing import Any

from q2_baseline.time_axis import date_range
from q3_baseline.exporter import export_result3_candidate

from .config import Q43Config


def export_result4_3_candidate(config: Q43Config, candidate_path: Path, plans: dict[date, object], executed: tuple[object, ...]) -> dict[str, Any]:
    official = config.official_result4_3_template.resolve()
    if candidate_path.resolve() == official:
        raise ValueError("candidate path must differ from official result4-3.xlsx")
    adapter = replace(config.q3, official_result3_template=official)
    january_end = date.fromordinal(config.output_start.toordinal() - 1)
    padded: dict[date, object] = {day: None for day in date_range(adapter.warmup_start, january_end)}
    padded.update(plans)
    info = export_result3_candidate(adapter, candidate_path, padded, executed)
    info["official_template_role"] = "result4-3.xlsx"
    return info


def toy_export_roundtrip(path: Path) -> dict[str, object]:
    """A tiny synthetic workbook check; never reads or writes the official template."""
    from openpyxl import Workbook, load_workbook

    if path.exists():
        raise FileExistsError(path)
    workbook = Workbook()
    first = workbook.active
    first.title = "计划购电量"
    for name in ("调整购电量", "充放电量", "紧急购电量"):
        workbook.create_sheet(name)
    first.append(["日期", "0:10-0:20", "合计"])
    first.append([date(2025, 2, 1), 1.25, 1.25])
    workbook.save(path); workbook.close()
    check = load_workbook(path, read_only=True, data_only=False)
    result = {
        "sheets": check.sheetnames,
        "value": check["计划购电量"].cell(2, 2).value,
        "passed": check.sheetnames == ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
        and check["计划购电量"].cell(2, 2).value == 1.25,
    }
    check.close()
    return result
