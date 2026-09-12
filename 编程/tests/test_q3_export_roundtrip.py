import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook

from q1_baseline.run_manifest import sha256_file
from q2_baseline.time_axis import target_day
from q3_baseline.config import Q3Config, Q3Parameters
from q3_baseline.export_validation import validate_result3_candidate
from q3_baseline.exporter import export_result3_candidate
from q3_baseline.state_machine import ExecutedInterval
from tests.helpers_q3 import zero_vector_plan

class Q3ExportRoundtripTests(unittest.TestCase):
    def test_full_four_sheet_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);template=root/"result3.xlsx";book=Workbook()
            plan=book.active;plan.title="计划购电量"
            for name in ("计划购电量","调整购电量"):
                sheet=plan if name=="计划购电量" else book.create_sheet(name)
                sheet.cell(1,1).value="日期\\时间"
                for slot in range(1,145):
                    sheet.cell(1,slot+1).value="0:10-0:20" if slot==1 else "0:00-0:10+1" if slot==144 else f"slot-{slot}"
                sheet.cell(1,146).value="全天购电量";sheet.cell(1,147).value="全天购电费"
                for i in range(334):sheet.cell(i+2,1).value=date(2025,2,1)+timedelta(days=i)
            storage=book.create_sheet("充放电量")
            for c,h in enumerate(("日期","时间段","充电量","放电量","时刻","储电量"),1):storage.cell(1,c).value=h
            storage.cell(2,1).value=date(2025,2,1)
            emergency=book.create_sheet("紧急购电量")
            for c,h in enumerate(("日期","购电时间段","购电量"),1):emergency.cell(1,c).value=h
            emergency.cell(2,1).value="⁝";book.save(template);book.close()
            config=Q3Config("toy","M2-Q2-EXPERIMENT-weekday_buffer-v1.0","toy",root/"price",root/"actual",root/"point",root/"zoh",root/"interp",root/"mapping",template,root/"manifest",root/"summary",root/"q2","highs",date(2025,1,1),date(2025,2,1),date(2025,12,31),{"0only":(0,),"0_12":(0,720),"rolling4":(0,360,720,1080)},Q3Parameters())
            plans={date(2025,1,1)+timedelta(days=i):zero_vector_plan(date(2025,1,1)+timedelta(days=i)) for i in range(365)}
            executed=[];soc=6000.
            for day,rolling in plans.items():
                for target in target_day(day):
                    c=2.0 if target.interval_start.isoformat()=="2025-02-01T00:00:00" else 0.0
                    d=1.62 if target.interval_start.isoformat()=="2025-02-01T00:10:00" else 0.0
                    e=1.0 if target.interval_start.isoformat()=="2025-02-01T00:00:00" else 2.0 if target.interval_start.isoformat()=="2025-02-01T00:10:00" else 0.0
                    before=soc;soc=soc+.9*c-d/.9
                    executed.append(ExecutedInterval(day.isoformat(),target.template_slot,target.interval_start,target.interval_end,10.,10.,c,d,e,0.,before,soc,1.,"attachment3",None,0.,0.,10.,10.,0.,0.,0.,10.,5*e,10.+5*e,"MODEL_B"))
            candidate=root/"result3_candidate.xlsx";original=sha256_file(template)
            export_result3_candidate(config,candidate,plans,tuple(executed))
            report=validate_result3_candidate(config,candidate,plans,tuple(executed),original)
            self.assertTrue(report["passed"],report)
            result=load_workbook(candidate,read_only=True,data_only=True)
            self.assertEqual(result["紧急购电量"].cell(2,3).value,3.0)
            result.close();self.assertEqual(sha256_file(template),original)

if __name__=="__main__":unittest.main()
