from __future__ import annotations

import re
import tempfile
import unittest
from dataclasses import replace
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from q1_baseline.run_manifest import sha256_file
from q2_baseline.config import load_config as load_q2_config
from q2_baseline.data import read_q2_inputs
from q2_baseline.planner import build_daily_milp
from q2_baseline.time_axis import target_day
from q4_baseline.config import PREDICTORS, Q4Config, Q4Parameters, load_config
from q4_baseline.export_validation import validate_result4_2_candidate
from q4_baseline.exporter import export_result4_2_candidate
from q4_baseline.forecast import build_price_forecast_day
from q4_baseline.integrity import verify_q4_inputs
from q4_baseline.price import PriceHistory, PriceRecord, read_price_records
from q4_baseline.reference import FrozenQ2Reference
from q4_baseline.runner import _assert_identity_unchanged, _identity_snapshot, price_unaware_resettlement
from q4_baseline.settlement import Q4ReplayDay
from tests.helpers_q2 import zero_replay


class Q4StaticGateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.config = load_config(Path("config/q4_baseline.json").resolve())

    def test_frozen_template_boundaries_include_slot144_next_day_interval(self) -> None:
        jan31 = target_day(date(2025, 1, 31))
        feb1 = target_day(date(2025, 2, 1))
        dec31 = target_day(date(2025, 12, 31))

        self.assertEqual((jan31[143].interval_start, jan31[143].interval_end), (datetime(2025, 2, 1, 0, 0), datetime(2025, 2, 1, 0, 10)))
        self.assertEqual((feb1[0].interval_start, feb1[0].interval_end), (datetime(2025, 2, 1, 0, 10), datetime(2025, 2, 1, 0, 20)))
        self.assertEqual((feb1[142].interval_start, feb1[142].interval_end), (datetime(2025, 2, 1, 23, 50), datetime(2025, 2, 2, 0, 0)))
        self.assertEqual((feb1[143].interval_start, feb1[143].interval_end), (datetime(2025, 2, 2, 0, 0), datetime(2025, 2, 2, 0, 10)))
        self.assertEqual((dec31[143].interval_start, dec31[143].interval_end), (datetime(2026, 1, 1, 0, 0), datetime(2026, 1, 1, 0, 10)))

    def test_q4_price_forecast_uses_the_same_slot_boundaries(self) -> None:
        history = PriceHistory(read_price_records(self.config.normalized_price_input), self.config.parameters)
        forecast = build_price_forecast_day(date(2025, 12, 31), "P1_WEEKDAY_SAME_CLOCK", history)
        self.assertEqual(forecast.rows[0]["interval_start"], datetime(2025, 12, 31, 0, 10))
        self.assertEqual(forecast.rows[-1]["interval_start"], datetime(2026, 1, 1, 0, 0))
        self.assertEqual(forecast.rows[-1]["interval_end"], datetime(2026, 1, 1, 0, 10))

    def test_official_template_hash_is_full_and_enforced(self) -> None:
        expected = self.config.official_result4_2_sha256
        self.assertRegex(expected, re.compile(r"^[0-9A-F]{64}$"))
        identity = verify_q4_inputs(self.config)
        self.assertEqual(identity["sha256"]["official_result4_2"], expected)
        with self.assertRaisesRegex(ValueError, "SHA-256"):
            replace(self.config, official_result4_2_sha256="G" * 64).validate()
        with self.assertRaisesRegex(RuntimeError, "Q4_RESULT4_2_TEMPLATE_HASH_HARD_FAIL"):
            verify_q4_inputs(replace(self.config, official_result4_2_sha256="0" * 64))

    def test_future_price_perturbation_does_not_change_any_predictor_plan(self) -> None:
        start = datetime(2025, 1, 1)
        rows = tuple(
            PriceRecord((start + timedelta(days=day)).date(), slot + 1, start + timedelta(days=day, minutes=10 * slot), start + timedelta(days=day, minutes=10 * (slot + 1)), 0.2 + 0.001 * slot + 0.01 * day)
            for day in range(21)
            for slot in range(144)
        )
        decision = datetime(2025, 1, 21)
        perturbed = tuple(replace(row, price=row.price * 100) if row.interval_start > decision else row for row in rows)
        for predictor in PREDICTORS:
            original = build_price_forecast_day(decision.date(), predictor, PriceHistory(rows, self.config.parameters))
            changed = build_price_forecast_day(decision.date(), predictor, PriceHistory(perturbed, self.config.parameters))
            self.assertEqual([row["price_plan"] for row in original.rows], [row["price_plan"] for row in changed.rows])

    def test_each_forecast_day_builds_one_causal_view(self) -> None:
        class CountingPriceHistory(PriceHistory):
            def __init__(self, records, params):
                super().__init__(records, params)
                self.view_calls = 0

            def view(self, decision_time):
                self.view_calls += 1
                return super().view(decision_time)

        rows = read_price_records(self.config.normalized_price_input)
        for predictor in PREDICTORS:
            history = CountingPriceHistory(rows, self.config.parameters)
            build_price_forecast_day(date(2025, 2, 1), predictor, history)
            self.assertEqual(history.view_calls, 1, predictor)

    def test_human_run_instructions_list_all_tracks_in_required_order(self) -> None:
        text = Path("Q4-2_人工运行说明.md").read_text(encoding="utf-8")
        sequence = (
            "--track REF_Q2_FIXED_PRICE_FROZEN",
            "--track Q4_2_PRICE_UNAWARE_RESETTLEMENT",
            "--track Q4_2_PRICE_BASELINE_P0",
            "--track Q4_2_PRICE_CANDIDATE_P1",
            "--track Q4_2_PRICE_CANDIDATE_P2",
            "--track Q4_2_PRICE_CANDIDATE_P3",
            "--track Q4_2_NOSTORAGE",
            " compare ",
        )
        positions = [text.index(item) for item in sequence]
        self.assertEqual(positions, sorted(positions))
        self.assertIn('--predictor "$selectedPredictor"', text)
        self.assertIn('--selected-predictor "$selectedTrack"', text)
        self.assertIn("历史固定价锚点", text)
        self.assertIn("price-awareness 价值的主参考", text)

    def test_all_predictors_share_frozen_feb1_initial_soc(self) -> None:
        q2_config = load_q2_config(self.config.q2_config)
        data = read_q2_inputs(q2_config.normalized_price_input, self.config.normalized_actual_input, q2_config.parameters)
        reference = FrozenQ2Reference(self.config.q2_reference_run / "dispatch_timeseries.csv", data)
        history = PriceHistory(read_price_records(self.config.normalized_price_input), self.config.parameters)
        initial_soc = reference.initial_soc_feb1()
        self.assertEqual(initial_soc, 1200.0)
        for predictor in PREDICTORS:
            price_day = build_price_forecast_day(date(2025, 2, 1), predictor, history)
            problem = build_daily_milp(reference.forecast(date(2025, 2, 1)), price_day.price_plan, self.config.parameters, initial_soc)
            row = problem.row_names.index("soc_window_start_0010")
            self.assertEqual(problem.row_lower[row], initial_soc)
            self.assertEqual(problem.row_upper[row], initial_soc)

    def test_price_unaware_resettlement_preserves_frozen_q2_actions(self) -> None:
        q2_config = load_q2_config(self.config.q2_config)
        data = read_q2_inputs(q2_config.normalized_price_input, self.config.normalized_actual_input, q2_config.parameters)
        reference = FrozenQ2Reference(self.config.q2_reference_run / "dispatch_timeseries.csv", data)
        history = PriceHistory(read_price_records(self.config.normalized_price_input), self.config.parameters)
        day = date(2025, 2, 1)
        frozen = reference.plan(day)
        replay = price_unaware_resettlement(reference, history, day, self.config)
        for actual, unchanged in ((replay.plan.G, frozen.G), (replay.plan.C, frozen.C), (replay.plan.D, frozen.D), (replay.plan.S, frozen.S)):
            np.testing.assert_array_equal(actual, unchanged)
        np.testing.assert_array_equal(replay.price_actual, np.asarray([history.by_start[row.interval_start].price for row in replay.actual]))

    def test_identity_gate_rejects_config_change_during_run(self) -> None:
        input_identity = verify_q4_inputs(self.config)
        with tempfile.TemporaryDirectory() as directory:
            config_copy = Path(directory) / "q4.json"
            config_copy.write_bytes(Path("config/q4_baseline.json").read_bytes())
            initial = _identity_snapshot(self.config, config_copy, input_identity, 1200.0)
            config_copy.write_text(config_copy.read_text(encoding="utf-8") + "\n", encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "Q4_IDENTITY_HARD_FAIL"):
                _assert_identity_unchanged(self.config, config_copy, initial, input_identity)

    def test_q4_exporter_toy_roundtrip_uses_locked_template_hash(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "result4-2.xlsx"
            workbook = Workbook()
            plan = workbook.active
            plan.title = "计划购电量"
            plan.cell(1, 1).value = "日期\\时间"
            for slot in range(1, 145):
                plan.cell(1, slot + 1).value = "0:10-0:20" if slot == 1 else "0:00-0:10+1" if slot == 144 else f"slot-{slot}"
            plan.cell(1, 146).value = "全天购电量"
            plan.cell(1, 147).value = "全天购电费"
            for index in range(334):
                plan.cell(index + 2, 1).value = date(2025, 2, 1) + timedelta(days=index)
            storage = workbook.create_sheet("充放电量")
            for column, header in enumerate(("日期", "时间段", "充电量", "放电量", "时刻", "储电量"), start=1):
                storage.cell(1, column).value = header
            storage.cell(2, 1).value = date(2025, 2, 1)
            emergency = workbook.create_sheet("紧急购电量")
            for column, header in enumerate(("日期", "购电时间段", "购电量"), start=1):
                emergency.cell(1, column).value = header
            emergency.cell(2, 1).value = "⁝"
            workbook.save(template)
            workbook.close()

            template_hash = sha256_file(template)
            config = Q4Config(
                model_version="toy", q2_frozen_model_version="M2-Q2-EXPERIMENT-weekday_buffer-v1.0",
                q2_reference_run=root, data_version="toy", normalized_price_input=root / "price.csv",
                normalized_price_sha256="0" * 64, normalized_actual_input=root / "actual.csv", normalized_actual_sha256="0" * 64,
                q2_price_input=root / "q2_price.csv", q2_config=root / "q2.json", q2_reference_dispatch_sha256="0" * 64,
                official_result4_2_template=template, official_result4_2_sha256=template_hash,
                audit_manifest=root / "manifest.json", audit_summary=root / "audit.json", solver_name="highs",
                warmup_start=date(2025, 1, 1), output_start=date(2025, 2, 1), output_end=date(2025, 12, 31), parameters=Q4Parameters(),
            )
            q4_replays = []
            for index in range(365):
                replay = zero_replay(date(2025, 1, 1) + timedelta(days=index))
                q4_replays.append(Q4ReplayDay(replay.plan, replay.actual, np.ones(144), replay.emergency_kwh, replay.surplus_kwh, replay.planned_purchase_cost, replay.emergency_purchase_cost))
            candidate = root / "toy_result4-2_candidate.xlsx"
            export_result4_2_candidate(config, candidate, tuple(q4_replays))
            report = validate_result4_2_candidate(config, candidate, tuple(q4_replays))
            self.assertTrue(report["passed"], report)


if __name__ == "__main__":
    unittest.main()
