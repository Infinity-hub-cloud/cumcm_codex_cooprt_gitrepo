from __future__ import annotations

import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from q1_baseline.run_manifest import sha256_file

from q2_baseline.config import Q2Config, Q2Parameters
from q2_baseline.export_validation import validate_result2_candidate
from q2_baseline.exporter import export_result2_candidate

from tests.helpers_q2 import zero_replay


class Q2ExportRoundtripTests(unittest.TestCase):
    def test_toy_full_calendar_export_and_reread(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "result2.xlsx"
            workbook = Workbook()
            plan = workbook.active
            plan.title = "计划购电量"
            plan.cell(1, 1).value = "日期\\时间"
            for slot in range(1, 145):
                plan.cell(1, slot + 1).value = (
                    "0:10-0:20" if slot == 1 else
                    "0:00-0:10+1" if slot == 144 else f"slot-{slot}"
                )
            plan.cell(1, 146).value = "全天购电量"
            plan.cell(1, 147).value = "全天购电费"
            for index in range(334):
                plan.cell(index + 2, 1).value = date(2025, 2, 1) + timedelta(days=index)
            storage = workbook.create_sheet("充放电量")
            for column, header in enumerate(("日期", "时间段", "充电量", "放电量", "时刻", "储电量"), start=1):
                storage.cell(1, column).value = header
            storage.cell(2, 1).value = date(2025, 2, 1)
            emergency = workbook.create_sheet("紧急购电量")
            for column, header in enumerate(("日期", "购电时间段", "购电量"), start=1):
                emergency.cell(1, column).value = header
            emergency.cell(2, 1).value = "⁝"
            workbook.save(template)
            workbook.close()

            config = Q2Config(
                model_version="toy",
                forecast_version="toy",
                data_version="toy",
                normalized_price_input=root / "price.csv",
                normalized_actual_input=root / "actual.csv",
                official_result2_template=template,
                audit_manifest=root / "manifest.json",
                audit_summary=root / "audit_summary.json",
                solver_name="highs",
                warmup_start=date(2025, 1, 1),
                output_start=date(2025, 2, 1),
                output_end=date(2025, 12, 31),
                parameters=Q2Parameters(),
            )
            def replay_for(day: date):
                if day == date(2025, 1, 31):
                    return zero_replay(day, {143: 1.0})
                if day == date(2025, 2, 1):
                    return zero_replay(day, {0: 2.0, 2: 3.0})
                return zero_replay(day)

            replays = tuple(
                replay_for(date(2025, 1, 1) + timedelta(days=index))
                for index in range(365)
            )
            candidate = root / "result2_candidate.xlsx"
            template_hash = sha256_file(template)
            export_result2_candidate(config, candidate, replays)
            report = validate_result2_candidate(config, candidate, replays, template_hash)
            self.assertTrue(report["passed"], report)
            result = load_workbook(candidate, read_only=True, data_only=True)
            emergency_result = result["紧急购电量"]
            self.assertEqual(emergency_result.max_row, 3)
            self.assertIsNotNone(emergency_result.cell(2, 1).value)
            self.assertIsNone(emergency_result.cell(3, 1).value)
            result.close()


if __name__ == "__main__":
    unittest.main()
