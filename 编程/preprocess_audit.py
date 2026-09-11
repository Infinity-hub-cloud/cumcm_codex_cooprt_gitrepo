from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook


DEFAULT_INTERVAL_COUNT = 144
DEFAULT_DELTA_MINUTES = 10
DEFAULT_HOURLY_STEP_MINUTES = 60
DEFAULT_HOURLY_HORIZON = 24
DEFAULT_START_DATE = dt.date(2025, 1, 1)
DEFAULT_END_DATE = dt.date(2025, 12, 31)
DEFAULT_OUTPUT_START_DATE = dt.date(2025, 2, 1)
DEFAULT_OUTPUT_END_DATE = dt.date(2025, 12, 31)
DEFAULT_UPDATE_TIMES = ("0:00", "6:00", "12:00", "18:00")
MAX_ISSUE_DETAILS = 200


@dataclass(frozen=True)
class AuditConfig:
    interval_count: int = DEFAULT_INTERVAL_COUNT
    delta_minutes: int = DEFAULT_DELTA_MINUTES
    hourly_step_minutes: int = DEFAULT_HOURLY_STEP_MINUTES
    hourly_horizon: int = DEFAULT_HOURLY_HORIZON
    start_date: dt.date = DEFAULT_START_DATE
    end_date: dt.date = DEFAULT_END_DATE
    output_start_date: dt.date = DEFAULT_OUTPUT_START_DATE
    output_end_date: dt.date = DEFAULT_OUTPUT_END_DATE
    update_times: tuple[str, ...] = DEFAULT_UPDATE_TIMES

    @property
    def expected_sample_minutes(self) -> tuple[int, ...]:
        return expected_sample_minutes(self.interval_count, self.delta_minutes)

    @property
    def delta_hours(self) -> float:
        return self.delta_minutes / 60.0


@dataclass(frozen=True)
class MappedHourlyTarget:
    target_date: dt.date
    target_start_minute: int
    target_end_minute: int
    interval_indices: tuple[int, ...]


@dataclass(frozen=True)
class MappedInstantaneousSample:
    interval_start_date: dt.date
    interval_start_minute: int
    interval_end_date: dt.date
    interval_end_minute: int


class AuditLog:
    def __init__(self) -> None:
        self.checks: list[dict[str, Any]] = []
        self.errors: list[dict[str, Any]] = []
        self.warnings: list[dict[str, Any]] = []
        self.stats: dict[str, Any] = {}

    def check(self, name: str, passed: bool, detail: str, severity: str = "error") -> None:
        record = {"name": name, "passed": bool(passed), "detail": detail, "severity": severity}
        self.checks.append(record)
        if passed:
            return
        target = self.errors if severity == "error" else self.warnings
        if len(target) < MAX_ISSUE_DETAILS:
            target.append(record)

    def error(self, name: str, detail: str) -> None:
        self.check(name, False, detail, "error")

    def warning(self, name: str, detail: str) -> None:
        self.check(name, False, detail, "warning")

    def status(self) -> str:
        if self.errors:
            return "FAIL"
        if self.warnings:
            return "WARN"
        return "PASS"


def parse_date_value(value: Any) -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isfinite(float(value)) and 1 <= float(value) <= 100000:
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))).date()
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_clock_minutes(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, dt.time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if 0 <= number <= 1:
            return int(round(number * 1440))
        if number.is_integer() and 0 <= number <= 1440:
            return int(number)
        return None
    text = str(value).strip().replace(" ", "")
    match = re.fullmatch(r"(\d{1,2}):(\d{1,2})(?::(\d{2}))?(\+1)?", text)
    if not match:
        return None
    hour = int(match.group(1))
    minute = int(match.group(2))
    second = int(match.group(3) or 0)
    next_day = bool(match.group(4))
    if minute >= 60 or second >= 60:
        return None
    if hour == 24 and minute == 0 and second == 0 and not next_day:
        return 1440
    if hour >= 24:
        return None
    result = hour * 60 + minute
    if second >= 30:
        result += 1
    if next_day:
        result += 1440
    return result if 0 <= result <= 2880 else None


def parse_interval_label(value: Any) -> tuple[int, int] | None:
    if value is None:
        return None
    text = str(value).strip()
    parts = text.split("-", 1)
    if len(parts) != 2:
        return None
    start = parse_clock_minutes(parts[0])
    end = parse_clock_minutes(parts[1])
    if start is None or end is None:
        return None
    if end > 1440 and start < end - 1440:
        start += 1440
    if end <= start:
        end += 1440
    return start, end


def expected_sample_minutes(interval_count: int, delta_minutes: int) -> tuple[int, ...]:
    return tuple((index + 1) * delta_minutes for index in range(interval_count))


def map_instantaneous_sample(
    sample_date: dt.date,
    sample_minutes: int,
    delta_minutes: int,
) -> MappedInstantaneousSample:
    start_day_offset, start_minute = divmod(sample_minutes, 1440)
    absolute_end = sample_minutes + delta_minutes
    end_day_offset, end_minute = divmod(absolute_end, 1440)
    return MappedInstantaneousSample(
        interval_start_date=sample_date + dt.timedelta(days=start_day_offset),
        interval_start_minute=start_minute,
        interval_end_date=sample_date + dt.timedelta(days=end_day_offset),
        interval_end_minute=end_minute,
    )


def map_hourly_target(
    issue_date: dt.date,
    issue_minutes: int,
    lead_hour: int,
    delta_minutes: int,
    hourly_step_minutes: int,
) -> MappedHourlyTarget:
    if lead_hour < 1:
        raise ValueError("lead_hour must be positive")
    if hourly_step_minutes % delta_minutes != 0:
        raise ValueError("hourly step must be divisible by delta minutes")
    absolute_start = issue_minutes + (lead_hour - 1) * hourly_step_minutes
    day_offset, start_minute = divmod(absolute_start, 1440)
    end_minute = start_minute + hourly_step_minutes
    if end_minute > 1440:
        raise ValueError("hourly target cannot span more than one calendar day")
    slot_count = hourly_step_minutes // delta_minutes
    first_index = start_minute // delta_minutes + 1
    indices = tuple(first_index + offset for offset in range(slot_count))
    return MappedHourlyTarget(
        target_date=issue_date + dt.timedelta(days=day_offset),
        target_start_minute=start_minute,
        target_end_minute=end_minute,
        interval_indices=indices,
    )


def safe_float(value: Any) -> tuple[float | None, str | None]:
    if value is None or value == "":
        return None, "missing"
    if isinstance(value, bool):
        return None, "boolean"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None, "non_numeric"
    if not math.isfinite(number):
        return None, "non_finite"
    return number, None


def fmt_number(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.4f}"


def fmt_date(value: dt.date | None) -> str:
    return "" if value is None else value.isoformat()


def row_values(sheet: Any) -> list[list[Any]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def write_csv(path: Path, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def record_numeric_issue(log: AuditLog, name: str, reason: str, location: str) -> None:
    if reason == "missing":
        log.error(name, f"缺失值：{location}")
    elif reason == "non_finite":
        log.error(name, f"非有限值：{location}")
    else:
        log.error(name, f"{reason}：{location}")


def numeric_stats(values: Iterable[float]) -> dict[str, Any]:
    numbers = list(values)
    if not numbers:
        return {"count": 0, "min": None, "max": None, "mean": None}
    return {
        "count": len(numbers),
        "min": min(numbers),
        "max": max(numbers),
        "mean": sum(numbers) / len(numbers),
    }


def audit_attachment1(path: Path, output_dir: Path, config: AuditConfig, log: AuditLog) -> None:
    name = "附件1.xlsx"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        log.error("attachment1_open", f"无法读取 {name}: {exc}")
        return
    log.check("attachment1_sheet_names", wb.sheetnames == ["Sheet1"], f"实际工作表：{wb.sheetnames}")
    sheet = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    rows = row_values(sheet)
    log.check("attachment1_dimensions", len(rows) == config.interval_count + 1 and sheet.max_column == 4, f"rows={len(rows)}, cols={sheet.max_column}")
    if not rows:
        wb.close()
        return
    headers = rows[0][:4]
    log.check("attachment1_headers", headers == ["时间", "电价", "小区负载", "光伏发电预测功率"], f"headers={headers}")
    sample_minutes_values: list[int | None] = []
    normalized: list[list[Any]] = []
    prices: list[float] = []
    loads: list[float] = []
    pv_values: list[float] = []
    for row_index, row in enumerate(rows[1:], start=2):
        time_value = row[0] if len(row) > 0 else None
        sample_minutes = parse_clock_minutes(time_value)
        sample_minutes_values.append(sample_minutes)
        if sample_minutes is None:
            log.error("attachment1_time_parse", f"行 {row_index} 时间无法解析：{time_value!r}")
        mapped = map_instantaneous_sample(dt.date(2000, 1, 1), sample_minutes, config.delta_minutes) if sample_minutes is not None else None
        record: list[float | None] = []
        for column_index, column_name in enumerate(("price", "load", "pv"), start=1):
            value, reason = safe_float(row[column_index] if len(row) > column_index else None)
            record.append(value)
            if reason:
                record_numeric_issue(log, f"attachment1_{column_name}", reason, f"行 {row_index} 列 {column_index + 1}")
            elif value is not None:
                if column_name == "price":
                    prices.append(value)
                elif column_name == "load":
                    loads.append(value)
                else:
                    pv_values.append(value)
                if value < 0:
                    log.error("attachment1_negative_value", f"行 {row_index} {column_name}={value}")
        price, load, pv = record
        normalized.append([
            row_index - 1,
            str(time_value) if time_value is not None else "",
            sample_minutes if sample_minutes is not None else "",
            (mapped.interval_start_date - dt.date(2000, 1, 1)).days if mapped else "",
            mapped.interval_start_minute if mapped else "",
            (mapped.interval_end_date - dt.date(2000, 1, 1)).days if mapped else "",
            mapped.interval_end_minute if mapped else "",
            fmt_number(price),
            fmt_number(load),
            fmt_number(pv),
            fmt_number(load * config.delta_hours if load is not None else None),
            fmt_number(pv * config.delta_hours if pv is not None else None),
        ])
    expected = config.expected_sample_minutes
    log.check("attachment1_time_axis", tuple(sample_minutes_values) == expected, f"actual_first_last={sample_minutes_values[:1]}...{sample_minutes_values[-1:]}, expected_first_last={expected[:1]}...{expected[-1:]}")
    log.stats[name] = {"dimensions": [len(rows), sheet.max_column], "price": numeric_stats(prices), "load_kw": numeric_stats(loads), "pv_kw": numeric_stats(pv_values)}
    write_csv(output_dir / "normalized" / "attachment1_single_day.csv", ["interval_index", "sample_time_label", "sample_minute", "interval_start_day_offset", "interval_start_minute", "interval_end_day_offset", "interval_end_minute", "price_yuan_per_kwh", "load_kw", "pv_forecast_kw", "load_kwh", "pv_kwh"], normalized)
    wb.close()


def audit_matrix_sheet(sheet: Any, sheet_name: str, config: AuditConfig, log: AuditLog, nonnegative: bool) -> tuple[dict[dt.date, list[float | None]], list[int | None], dict[str, Any]]:
    rows = row_values(sheet)
    log.check(f"{sheet_name}_dimensions", len(rows) == config.end_date.toordinal() - config.start_date.toordinal() + 2 and sheet.max_column == config.interval_count + 1, f"rows={len(rows)}, cols={sheet.max_column}")
    if not rows:
        return {}, [], {}
    header = rows[0]
    endpoints = [parse_clock_minutes(value) for value in header[1:]]
    log.check(f"{sheet_name}_time_axis", tuple(endpoints) == config.expected_sample_minutes, f"actual_first_last={endpoints[:1]}...{endpoints[-1:]}")
    values_by_date: dict[dt.date, list[float | None]] = {}
    dates: list[dt.date] = []
    numeric_values: list[float] = []
    for row_index, row in enumerate(rows[1:], start=2):
        date_value = parse_date_value(row[0] if row else None)
        if date_value is None:
            log.error(f"{sheet_name}_date_parse", f"行 {row_index} 日期无法解析：{row[0] if row else None!r}")
            continue
        dates.append(date_value)
        if date_value in values_by_date:
            log.error(f"{sheet_name}_duplicate_date", f"重复日期 {date_value}")
        row_values_float: list[float | None] = []
        for column_index in range(config.interval_count):
            raw = row[column_index + 1] if len(row) > column_index + 1 else None
            value, reason = safe_float(raw)
            row_values_float.append(value)
            if reason:
                record_numeric_issue(log, f"{sheet_name}_numeric", reason, f"日期 {date_value} 列 {column_index + 2}")
            elif value is not None:
                numeric_values.append(value)
                if nonnegative and value < 0:
                    log.error(f"{sheet_name}_negative_value", f"日期 {date_value} 列 {column_index + 2}={value}")
        values_by_date[date_value] = row_values_float
    expected_dates = [config.start_date + dt.timedelta(days=offset) for offset in range((config.end_date - config.start_date).days + 1)]
    log.check(f"{sheet_name}_date_coverage", dates == expected_dates, f"count={len(dates)}, first_last={dates[:1]}...{dates[-1:]}")
    log.check(f"{sheet_name}_duplicate_dates", len(dates) == len(set(dates)), f"unique={len(set(dates))}, rows={len(dates)}")
    return values_by_date, endpoints, {"dimensions": [len(rows), sheet.max_column], "date_count": len(dates), "values": numeric_stats(numeric_values)}


def audit_attachment2(path: Path, output_dir: Path, config: AuditConfig, log: AuditLog) -> None:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        log.error("attachment2_open", f"无法读取附件2：{exc}")
        return
    expected_sheets = ["小区负载", "光伏发电实际功率"]
    log.check("attachment2_sheet_names", wb.sheetnames == expected_sheets, f"实际工作表：{wb.sheetnames}")
    load_data: dict[dt.date, list[float | None]] = {}
    pv_data: dict[dt.date, list[float | None]] = {}
    sheet_stats: dict[str, Any] = {}
    if "小区负载" in wb.sheetnames:
        load_data, _, sheet_stats["小区负载"] = audit_matrix_sheet(wb["小区负载"], "attachment2_load", config, log, True)
    if "光伏发电实际功率" in wb.sheetnames:
        pv_data, _, sheet_stats["光伏发电实际功率"] = audit_matrix_sheet(wb["光伏发电实际功率"], "attachment2_pv", config, log, True)
    all_dates = sorted(set(load_data) | set(pv_data))
    normalized: list[list[Any]] = []
    for date_value in all_dates:
        for index in range(config.interval_count):
            load = load_data.get(date_value, [None] * config.interval_count)[index]
            pv = pv_data.get(date_value, [None] * config.interval_count)[index]
            sample_minutes = (index + 1) * config.delta_minutes
            mapped = map_instantaneous_sample(date_value, sample_minutes, config.delta_minutes)
            normalized.append([fmt_date(date_value), index + 1, sample_minutes, fmt_date(mapped.interval_start_date), mapped.interval_start_minute, fmt_date(mapped.interval_end_date), mapped.interval_end_minute, fmt_number(load), fmt_number(pv), fmt_number(load * config.delta_hours if load is not None else None), fmt_number(pv * config.delta_hours if pv is not None else None)])
    log.stats["附件2.xlsx"] = {"sheets": sheet_stats, "joined_date_count": len(all_dates), "joined_row_count": len(normalized)}
    write_csv(output_dir / "normalized" / "attachment2_actual_long.csv", ["source_date", "interval_index", "sample_minute", "interval_start_date", "interval_start_minute", "interval_end_date", "interval_end_minute", "load_kw", "pv_actual_kw", "load_kwh", "pv_actual_kwh"], normalized)
    wb.close()


def audit_attachment4(path: Path, output_dir: Path, config: AuditConfig, log: AuditLog) -> None:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        log.error("attachment4_open", f"无法读取附件4：{exc}")
        return
    log.check("attachment4_sheet_names", wb.sheetnames == ["Sheet1"], f"实际工作表：{wb.sheetnames}")
    if "Sheet1" not in wb.sheetnames:
        wb.close()
        return
    price_data, _, sheet_stats = audit_matrix_sheet(wb["Sheet1"], "attachment4_price", config, log, False)
    negative_values = [value for row in price_data.values() for value in row if value is not None and value < 0]
    if negative_values:
        log.warning("PRICE-NEG-001", f"附件4发现 {len(negative_values)} 个负电价；在人工确认规则前不得直接求解")
    normalized: list[list[Any]] = []
    for date_value in sorted(price_data):
        for index, price in enumerate(price_data[date_value]):
            sample_minutes = (index + 1) * config.delta_minutes
            mapped = map_instantaneous_sample(date_value, sample_minutes, config.delta_minutes)
            normalized.append([fmt_date(date_value), index + 1, sample_minutes, fmt_date(mapped.interval_start_date), mapped.interval_start_minute, fmt_date(mapped.interval_end_date), mapped.interval_end_minute, fmt_number(price)])
    log.stats["附件4.xlsx"] = {"sheet": sheet_stats, "negative_price_count": len(negative_values), "row_count": len(normalized)}
    write_csv(output_dir / "normalized" / "attachment4_actual_price_long.csv", ["source_date", "interval_index", "sample_minute", "interval_start_date", "interval_start_minute", "interval_end_date", "interval_end_minute", "price_yuan_per_kwh"], normalized)
    wb.close()


def audit_attachment3(path: Path, output_dir: Path, config: AuditConfig, log: AuditLog) -> None:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        log.error("attachment3_open", f"无法读取附件3：{exc}")
        return
    log.check("attachment3_sheet_names", wb.sheetnames == ["Sheet1"], f"实际工作表：{wb.sheetnames}")
    if "Sheet1" not in wb.sheetnames:
        wb.close()
        return
    sheet = wb["Sheet1"]
    rows = row_values(sheet)
    expected_rows = (config.end_date - config.start_date).days + 1
    log.check("attachment3_dimensions", len(rows) == expected_rows * len(config.update_times) + 1 and sheet.max_column == config.hourly_horizon + 2, f"rows={len(rows)}, cols={sheet.max_column}")
    if not rows:
        wb.close()
        return
    header = rows[0]
    expected_header = ["日期", "预报时刻"] + [f"预报{index}小时" for index in range(1, config.hourly_horizon + 1)]
    log.check("attachment3_headers", header[: config.hourly_horizon + 2] == expected_header, f"headers={header[:config.hourly_horizon + 2]}")
    current_date: dt.date | None = None
    seen: set[tuple[dt.date, int]] = set()
    update_counts: dict[dt.date, int] = {}
    row_issue_dates: list[dt.date] = []
    row_issue_minutes: list[int] = []
    long_rows: list[list[Any]] = []
    mapped_rows: list[list[Any]] = []
    for row_index, row in enumerate(rows[1:], start=2):
        parsed_date = parse_date_value(row[0] if row else None)
        if parsed_date is not None:
            current_date = parsed_date
        if current_date is None:
            log.error("attachment3_date_forward_fill", f"行 {row_index} 在首次日期前为空")
            continue
        issue_minutes = parse_clock_minutes(row[1] if len(row) > 1 else None)
        if issue_minutes is None:
            log.error("attachment3_issue_time_parse", f"行 {row_index} 发布时间无法解析：{row[1] if len(row) > 1 else None!r}")
            continue
        row_issue_dates.append(current_date)
        row_issue_minutes.append(issue_minutes)
        issue_key = (current_date, issue_minutes)
        if issue_minutes not in {parse_clock_minutes(item) for item in config.update_times}:
            log.error("attachment3_issue_time_set", f"行 {row_index} 发布时间不在配置集合：{row[1]!r}")
        if issue_key in seen:
            log.error("attachment3_duplicate_issue", f"重复发布记录：{current_date} {row[1]}")
        seen.add(issue_key)
        update_counts[current_date] = update_counts.get(current_date, 0) + 1
        for lead_hour in range(1, config.hourly_horizon + 1):
            raw = row[lead_hour + 1] if len(row) > lead_hour + 1 else None
            forecast, reason = safe_float(raw)
            if reason:
                record_numeric_issue(log, "attachment3_forecast_numeric", reason, f"行 {row_index} 预报{lead_hour}小时")
            if forecast is not None and forecast < 0:
                log.error("attachment3_negative_value", f"行 {row_index} 预报{lead_hour}小时={forecast}")
            target = None
            if issue_minutes is not None:
                try:
                    target = map_hourly_target(current_date, issue_minutes, lead_hour, config.delta_minutes, config.hourly_step_minutes)
                except ValueError as exc:
                    log.error("attachment3_mapping", f"行 {row_index} 预报{lead_hour}小时映射失败：{exc}")
            long_rows.append([fmt_date(current_date), str(row[1]), lead_hour, fmt_number(forecast), fmt_date(target.target_date if target else None), target.target_start_minute if target else "", target.target_end_minute if target else "", ",".join(str(item) for item in target.interval_indices) if target else ""])
            if target:
                for interval_index in target.interval_indices:
                    start_minute = (interval_index - 1) * config.delta_minutes
                    mapped_rows.append([fmt_date(current_date), str(row[1]), lead_hour, fmt_number(forecast), fmt_date(target.target_date), interval_index, start_minute, start_minute + config.delta_minutes])
    expected_issue_minutes = {parse_clock_minutes(item) for item in config.update_times}
    expected_dates = [config.start_date + dt.timedelta(days=offset) for offset in range(expected_rows)]
    log.check("attachment3_date_coverage", sorted(update_counts) == expected_dates, f"actual_first_last={sorted(update_counts)[:1]}...{sorted(update_counts)[-1:]}, expected_first_last={expected_dates[:1]}...{expected_dates[-1:]}")
    expected_row_dates = [date_value for date_value in expected_dates for _ in config.update_times]
    expected_row_minutes = [parse_clock_minutes(item) for _ in expected_dates for item in config.update_times]
    log.check("attachment3_record_order", row_issue_dates == expected_row_dates and row_issue_minutes == expected_row_minutes, "发布记录未严格按日期和配置发布时间顺序排列")
    for date_value in [config.start_date + dt.timedelta(days=offset) for offset in range(expected_rows)]:
        log.check("attachment3_daily_update_count", update_counts.get(date_value, 0) == len(config.update_times), f"{date_value} 更新次数={update_counts.get(date_value, 0)}")
    log.stats["附件3.xlsx"] = {"dimensions": [len(rows), sheet.max_column], "issue_record_count": len(seen), "long_forecast_row_count": len(long_rows), "mapped_10min_row_count": len(mapped_rows), "configured_issue_minutes": sorted(item for item in expected_issue_minutes if item is not None)}
    write_csv(output_dir / "normalized" / "attachment3_forecast_long.csv", ["issue_date", "issue_time", "lead_hour", "forecast_kw", "target_date", "target_start_minute", "target_end_minute", "mapped_interval_indices"], long_rows)
    write_csv(output_dir / "normalized" / "attachment3_mapped_10min.csv", ["issue_date", "issue_time", "lead_hour", "forecast_kw", "target_date", "interval_index", "interval_start_minute", "interval_end_minute"], mapped_rows)
    wb.close()


def template_expected_dimensions(file_name: str) -> dict[str, tuple[int, int]]:
    if file_name == "result1.xlsx":
        return {"计划购电量": (145, 2), "充放电量": (7, 5)}
    if file_name in {"result2.xlsx", "result4-2.xlsx"}:
        return {"计划购电量": (335, 147), "充放电量": (20, 6), "紧急购电量": (11, 3)}
    return {"计划购电量": (335, 147), "调整购电量": (335, 147), "充放电量": (26, 6), "紧急购电量": (11, 3)}


def audit_templates(template_dir: Path, config: AuditConfig, log: AuditLog) -> None:
    template_stats: dict[str, Any] = {}
    expected_template_names = {"result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"}
    actual_template_names = {path.name for path in template_dir.glob("*.xlsx") if not path.name.startswith("~$")}
    log.check("template_file_set", actual_template_names == expected_template_names, f"actual={sorted(actual_template_names)}, expected={sorted(expected_template_names)}")
    for path in sorted(path for path in template_dir.glob("*.xlsx") if not path.name.startswith("~$")):
        file_name = path.name
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            log.error("template_open", f"无法读取 {file_name}：{exc}")
            continue
        expected = template_expected_dimensions(file_name)
        log.check(f"{file_name}_sheet_names", set(wb.sheetnames) == set(expected), f"实际工作表：{wb.sheetnames}")
        file_meta: dict[str, Any] = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = row_values(sheet)
            dims = [len(rows), sheet.max_column]
            file_meta[sheet_name] = {"dimensions": dims, "first_rows": [row[:10] for row in rows[:3]], "nonempty_cells": sum(1 for row in rows for value in row if value not in (None, ""))}
            if sheet_name in expected:
                log.check(f"{file_name}_{sheet_name}_dimensions", tuple(dims) == expected[sheet_name], f"actual={dims}, expected={expected[sheet_name]}")
            if sheet_name in {"计划购电量", "调整购电量"} and rows:
                if file_name == "result1.xlsx" and sheet_name == "计划购电量":
                    labels = [str(row[0]) for row in rows[1 : 1 + config.interval_count] if row and row[0] not in (None, "")]
                else:
                    labels = [str(value) for value in rows[0][1 : 1 + config.interval_count] if value not in (None, "")]
                if labels:
                    intervals = [parse_interval_label(label) for label in labels]
                    log.check(f"{file_name}_{sheet_name}_interval_labels_parse", all(interval is not None for interval in intervals), f"first={labels[0]}, last={labels[-1]}")
                    expected_intervals = [(config.delta_minutes + index * config.delta_minutes, config.delta_minutes + (index + 1) * config.delta_minutes) for index in range(config.interval_count)]
                    log.check(f"{file_name}_{sheet_name}_interval_sequence", intervals == expected_intervals, f"actual_first_last={intervals[:1]}...{intervals[-1:]}, expected_first_last={expected_intervals[:1]}...{expected_intervals[-1:]}")
                data_dates = [parse_date_value(row[0]) for row in rows[1:] if row and row[0] not in (None, "")]
                if data_dates and not (file_name == "result1.xlsx" and sheet_name == "计划购电量"):
                    log.check(f"{file_name}_{sheet_name}_date_start", data_dates[0] == config.output_start_date, f"first_date={data_dates[0]}")
                    log.check(f"{file_name}_{sheet_name}_date_end", data_dates[-1] == config.output_end_date, f"last_date={data_dates[-1]}")
        template_stats[file_name] = file_meta
        wb.close()
    log.stats["附件5_templates"] = template_stats


def audit_formula_cells(input_root: Path, log: AuditLog) -> None:
    formula_stats: dict[str, int] = {}
    for path in sorted(input_root.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        count = 0
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            count += 1
            wb.close()
        except Exception as exc:
            log.warning("formula_scan", f"{path.name} 公式扫描失败：{exc}")
            continue
        formula_stats[str(path.relative_to(input_root))] = count
        if count:
            log.warning("FORMULA-CELL-001", f"{path.relative_to(input_root)} 含 {count} 个公式单元格；请人工确认是否有缓存值")
    log.stats["formula_cells"] = formula_stats


def collect_manifest(input_root: Path) -> list[dict[str, Any]]:
    manifest: list[dict[str, Any]] = []
    for path in sorted(item for item in input_root.rglob("*") if item.is_file()):
        if path.name.startswith("~$"):
            continue
        stat = path.stat()
        manifest.append({"relative_path": str(path.relative_to(input_root)), "size_bytes": stat.st_size, "modified_time": dt.datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"), "sha256": sha256_file(path)})
    return manifest


def markdown_table(headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> str:
    lines = ["| " + " | ".join(str(item) for item in headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    lines.extend("| " + " | ".join(str(item).replace("|", "\\|") for item in row) + " |" for row in rows)
    return "\n".join(lines)


def write_reports(input_root: Path, output_dir: Path, config: AuditConfig, log: AuditLog, manifest: list[dict[str, Any]]) -> None:
    payload = {
        "status": log.status(),
        "generated_at": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "input_root": str(input_root),
        "output_dir": str(output_dir),
        "config": {**asdict(config), "start_date": config.start_date.isoformat(), "end_date": config.end_date.isoformat(), "output_start_date": config.output_start_date.isoformat(), "output_end_date": config.output_end_date.isoformat(), "update_times": list(config.update_times)},
        "checks": log.checks,
        "errors": log.errors,
        "warnings": log.warnings,
        "stats": log.stats,
        "manifest": manifest,
    }
    (output_dir / "audit").mkdir(parents=True, exist_ok=True)
    (output_dir / "audit" / "audit_summary.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    lines = [
        "# C 题数据预处理与人工审计报告",
        "",
        f"> 状态：`{log.status()}`（脚本只读运行；未生成模型、预测或求解结果）  ",
        f"> 生成时间：`{payload['generated_at']}`  ",
        f"> 输入目录：`{input_root}`  ",
        f"> 输出目录：`{output_dir}`  ",
        "",
        "## 1. 人工先看结论",
        "",
        f"- 错误数：`{len(log.errors)}`；告警数：`{len(log.warnings)}`。",
        "- 原始附件未被修改；规范化数据全部写入 `normalized/`。",
        "- 时间口径采用人工确认规则：附件 1/2/4 为瞬时时刻，按左端点保持映射到随后 10 分钟；模板标签严格保留。",
        "- 该报告不是模型运行结果；所有异常必须由人工逐项确认、关闭或登记。",
        "",
        "## 2. 参数接口快照",
        "",
        markdown_table(["参数", "值"], [["interval_count", config.interval_count], ["delta_minutes", config.delta_minutes], ["hourly_step_minutes", config.hourly_step_minutes], ["hourly_horizon", config.hourly_horizon], ["start_date", config.start_date], ["end_date", config.end_date], ["output_start_date", config.output_start_date], ["output_end_date", config.output_end_date], ["update_times", ", ".join(config.update_times)]]),
        "",
        "## 3. 检查结果",
        "",
        markdown_table(["检查", "状态", "级别", "详情"], [[item["name"], "PASS" if item["passed"] else "FAIL", item["severity"], item["detail"]] for item in log.checks]),
        "",
        "## 4. 错误与告警",
        "",
    ]
    if not log.errors and not log.warnings:
        lines.append("无。")
    else:
        for item in log.errors + log.warnings:
            lines.append(f"- `{item['severity'].upper()}` `{item['name']}`：{item['detail']}")
    lines.extend(["", "## 5. 工作簿统计", ""])
    for name, stat in log.stats.items():
        lines.append(f"### {name}")
        lines.append("")
        lines.append("```json")
        lines.append(json.dumps(stat, ensure_ascii=False, indent=2, default=str))
        lines.append("```")
        lines.append("")
    lines.extend(["## 6. 文件哈希清单", "", markdown_table(["相对路径", "字节数", "SHA-256"], [[item["relative_path"], item["size_bytes"], item["sha256"]] for item in manifest]), "", "## 7. 输出文件", "", "- `normalized/attachment1_single_day.csv`", "- `normalized/attachment2_actual_long.csv`", "- `normalized/attachment3_forecast_long.csv`", "- `normalized/attachment3_mapped_10min.csv`", "- `normalized/attachment4_actual_price_long.csv`", "- `audit/audit_summary.json`", "- 本报告 `audit/audit_report.md`", "", "## 8. 人工复核回传项", "", "1. 确认所有 `ERROR` 是否为真实异常、解析错误或可接受空值。", "2. 确认所有 `WARNING` 是否关闭、保留或升级为阻塞项。", "3. 抽查附件 1/2/4 的瞬时时刻是否映射到随后 10 分钟，特别是 0:00+1 到次日 0:00-0:10。", "4. 抽查附件 3 四个发布时间及跨日映射，并确认五个模板的 144 个时段标签顺序。", "5. 回传本报告、`audit_summary.json`、输入文件哈希和人工决定；在此之前不进入模型代码。", ""])
    (output_dir / "audit" / "audit_report.md").write_text("\n".join(lines), encoding="utf-8")
    (output_dir / "audit" / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="C题原始附件只读预处理与审计，不执行预测或优化")
    default_input = Path(__file__).resolve().parents[2] / "正式比赛C题"
    default_output = Path(__file__).resolve().parent / "预处理审计输出"
    parser.add_argument("--input-root", type=Path, default=default_input)
    parser.add_argument("--output-dir", type=Path, default=default_output)
    parser.add_argument("--interval-count", type=int, default=DEFAULT_INTERVAL_COUNT)
    parser.add_argument("--delta-minutes", type=int, default=DEFAULT_DELTA_MINUTES)
    parser.add_argument("--hourly-step-minutes", type=int, default=DEFAULT_HOURLY_STEP_MINUTES)
    parser.add_argument("--hourly-horizon", type=int, default=DEFAULT_HOURLY_HORIZON)
    parser.add_argument("--start-date", type=lambda value: dt.date.fromisoformat(value), default=DEFAULT_START_DATE)
    parser.add_argument("--end-date", type=lambda value: dt.date.fromisoformat(value), default=DEFAULT_END_DATE)
    parser.add_argument("--output-start-date", type=lambda value: dt.date.fromisoformat(value), default=DEFAULT_OUTPUT_START_DATE)
    parser.add_argument("--output-end-date", type=lambda value: dt.date.fromisoformat(value), default=DEFAULT_OUTPUT_END_DATE)
    parser.add_argument("--update-times", nargs="+", default=list(DEFAULT_UPDATE_TIMES))
    return parser


def run(args: argparse.Namespace) -> int:
    input_root = args.input_root.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    log = AuditLog()
    if not input_root.is_dir():
        print(f"输入目录不存在：{input_root}", file=sys.stderr)
        return 2
    if output_dir == input_root or output_dir.is_relative_to(input_root):
        print("输出目录不能位于原始输入目录内。", file=sys.stderr)
        return 2
    if output_dir.exists() and any(output_dir.iterdir()):
        print(f"输出目录非空，为避免覆盖请更换路径：{output_dir}", file=sys.stderr)
        return 2
    output_dir.mkdir(parents=True, exist_ok=False)
    config = AuditConfig(args.interval_count, args.delta_minutes, args.hourly_step_minutes, args.hourly_horizon, args.start_date, args.end_date, args.output_start_date, args.output_end_date, tuple(args.update_times))
    attachments = input_root / "附件"
    manifest = collect_manifest(input_root)
    required_files = [attachments / "附件1.xlsx", attachments / "附件2.xlsx", attachments / "附件3.xlsx", attachments / "附件4.xlsx", attachments / "附件5"]
    for path in required_files:
        log.check("required_path", path.exists(), f"{path.relative_to(input_root)}")
    if (attachments / "附件1.xlsx").is_file():
        audit_attachment1(attachments / "附件1.xlsx", output_dir, config, log)
    if (attachments / "附件2.xlsx").is_file():
        audit_attachment2(attachments / "附件2.xlsx", output_dir, config, log)
    if (attachments / "附件3.xlsx").is_file():
        audit_attachment3(attachments / "附件3.xlsx", output_dir, config, log)
    if (attachments / "附件4.xlsx").is_file():
        audit_attachment4(attachments / "附件4.xlsx", output_dir, config, log)
    if (attachments / "附件5").is_dir():
        audit_templates(attachments / "附件5", config, log)
    audit_formula_cells(input_root, log)
    write_reports(input_root, output_dir, config, log, manifest)
    print(f"审计完成：{log.status()}")
    print(f"报告：{output_dir / 'audit' / 'audit_report.md'}")
    print(f"规范化数据：{output_dir / 'normalized'}")
    return 2 if log.errors else 0


def main() -> int:
    return run(build_parser().parse_args())


if __name__ == "__main__":
    raise SystemExit(main())
