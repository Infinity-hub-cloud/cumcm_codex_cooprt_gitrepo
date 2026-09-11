from __future__ import annotations

import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from q1_baseline.run_manifest import sha256_file

from .config import Q2Config
from .replay import ReplayDay
from .time_axis import date_range


FOUR_HOUR_LABELS = (
    "0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00",
    "16:00-20:00", "20:00-24:00",
)


def round4(value: float) -> float:
    return float(Decimal(str(float(value))).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))


@dataclass(frozen=True)
class PhysicalDispatch:
    interval_start: datetime
    interval_end: datetime
    charge_kwh: float
    discharge_kwh: float
    emergency_kwh: float


@dataclass(frozen=True)
class EmergencyEvent:
    natural_date: date
    interval_start: datetime
    interval_end: datetime
    energy_kwh: float


def physical_dispatch_rows(replays: Iterable[ReplayDay]) -> tuple[PhysicalDispatch, ...]:
    rows: list[PhysicalDispatch] = []
    for replay in replays:
        for i, actual in enumerate(replay.actual):
            rows.append(
                PhysicalDispatch(
                    actual.interval_start,
                    actual.interval_end,
                    float(replay.plan.C[i]),
                    float(replay.plan.D[i]),
                    float(replay.emergency_kwh[i]),
                )
            )
    return tuple(sorted(rows, key=lambda row: row.interval_start))


def index_physical_by_natural_date(
    rows: Iterable[PhysicalDispatch],
) -> dict[date, tuple[PhysicalDispatch, ...]]:
    grouped: dict[date, list[PhysicalDispatch]] = {}
    for row in rows:
        grouped.setdefault(row.interval_start.date(), []).append(row)
    return {
        day: tuple(sorted(items, key=lambda item: item.interval_start))
        for day, items in grouped.items()
    }


def natural_day_four_hour_summary(
    rows: Iterable[PhysicalDispatch], natural_date: date
) -> tuple[tuple[float, float], ...]:
    selected = [row for row in rows if row.interval_start.date() == natural_date]
    if len(selected) != 144:
        raise ValueError(f"natural day {natural_date} must contain 144 physical intervals")
    selected.sort(key=lambda row: row.interval_start)
    expected = datetime.combine(natural_date, time.min)
    for i, row in enumerate(selected):
        if row.interval_start != expected + timedelta(minutes=10 * i):
            raise ValueError(f"natural-day physical interval gap at {natural_date} index {i}")
    return tuple(
        (
            float(sum(row.charge_kwh for row in selected[group * 24 : (group + 1) * 24])),
            float(sum(row.discharge_kwh for row in selected[group * 24 : (group + 1) * 24])),
        )
        for group in range(6)
    )


def merge_emergency_events(
    rows: Iterable[PhysicalDispatch], natural_date: date, tolerance: float
) -> tuple[EmergencyEvent, ...]:
    selected = sorted(
        (row for row in rows if row.interval_start.date() == natural_date),
        key=lambda row: row.interval_start,
    )
    if len(selected) != 144:
        raise ValueError(f"natural day {natural_date} must contain 144 physical intervals")
    events: list[EmergencyEvent] = []
    start: datetime | None = None
    end: datetime | None = None
    energy = 0.0
    for row in selected:
        if row.emergency_kwh > tolerance:
            if start is None:
                start = row.interval_start
                end = row.interval_end
                energy = row.emergency_kwh
            elif row.interval_start == end:
                end = row.interval_end
                energy += row.emergency_kwh
            else:
                events.append(EmergencyEvent(natural_date, start, end, energy))
                start, end, energy = row.interval_start, row.interval_end, row.emergency_kwh
        elif start is not None:
            events.append(EmergencyEvent(natural_date, start, end or start, energy))
            start = end = None
            energy = 0.0
    if start is not None:
        events.append(EmergencyEvent(natural_date, start, end or start, energy))
    return tuple(events)


def format_event_label(event: EmergencyEvent) -> str:
    midnight_next = datetime.combine(event.natural_date + timedelta(days=1), time.min)

    def label(value: datetime) -> str:
        if value == midnight_next:
            return "24:00"
        return f"{value.hour}:{value.minute:02d}"

    return f"{label(event.interval_start)}-{label(event.interval_end)}"


def _copy_row_style(sheet: Any, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def export_result2_candidate(
    config: Q2Config,
    candidate_path: Path,
    b1_replays: tuple[ReplayDay, ...],
) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for result2 export") from exc

    source = config.official_result2_template.resolve()
    candidate = candidate_path.resolve()
    if source == candidate:
        raise ValueError("candidate path must differ from official result2.xlsx")
    if candidate.exists():
        raise FileExistsError(f"refusing to overwrite candidate: {candidate}")
    original_hash = sha256_file(source)
    replay_by_date = {item.plan.template_date: item for item in b1_replays}
    all_dates = date_range(config.warmup_start, config.output_end)
    if tuple(sorted(replay_by_date)) != all_dates:
        raise ValueError("export requires complete Jan1-Dec31 replay chain")

    candidate.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, candidate)
    workbook = load_workbook(candidate, data_only=False)
    if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError(f"unexpected result2 sheets: {workbook.sheetnames}")

    output_dates = date_range(config.output_start, config.output_end)
    plan_sheet = workbook["计划购电量"]
    if (plan_sheet.max_row, plan_sheet.max_column) != (335, 147):
        raise ValueError("official plan sheet dimensions changed")
    if plan_sheet.cell(1, 2).value != "0:10-0:20" or plan_sheet.cell(1, 145).value != "0:00-0:10+1":
        raise ValueError("official template slot labels changed")
    for row_index, day in enumerate(output_dates, start=2):
        replay = replay_by_date[day]
        plan_sheet.cell(row_index, 1).value = day
        for slot in range(144):
            plan_sheet.cell(row_index, slot + 2).value = round4(replay.plan.G[slot])
            plan_sheet.cell(row_index, slot + 2).number_format = "0.0000"
        plan_sheet.cell(row_index, 146).value = round4(sum(replay.plan.G))
        plan_sheet.cell(row_index, 147).value = round4(sum(replay.planned_purchase_cost))
        plan_sheet.cell(row_index, 146).number_format = "0.0000"
        plan_sheet.cell(row_index, 147).number_format = "0.0000"

    physical = physical_dispatch_rows(b1_replays)
    physical_by_date = index_physical_by_natural_date(physical)
    storage = workbook["充放电量"]
    if storage.max_row > 1:
        storage.delete_rows(2, storage.max_row - 1)
    for row_index, day in enumerate(output_dates):
        summaries = natural_day_four_hour_summary(physical_by_date.get(day, ()), day)
        previous = replay_by_date[day - timedelta(days=1)].plan
        current = replay_by_date[day].plan
        soc_0000 = float(previous.S[-2])
        soc_2400 = float(current.S[-2])
        for group, label in enumerate(FOUR_HOUR_LABELS):
            target_row = 2 + row_index * 6 + group
            _copy_row_style(storage, 2, target_row, 6)
            storage.cell(target_row, 1).value = day if group == 0 else None
            storage.cell(target_row, 2).value = label
            storage.cell(target_row, 3).value = round4(summaries[group][0])
            storage.cell(target_row, 4).value = round4(summaries[group][1])
            storage.cell(target_row, 3).number_format = "0.0000"
            storage.cell(target_row, 4).number_format = "0.0000"
            if group == 0:
                storage.cell(target_row, 5).value = time(0, 0)
                storage.cell(target_row, 6).value = round4(soc_0000)
            elif group == 1:
                storage.cell(target_row, 5).value = "24:00"
                storage.cell(target_row, 6).value = round4(soc_2400)
            else:
                storage.cell(target_row, 5).value = None
                storage.cell(target_row, 6).value = None
            if group in (0, 1):
                storage.cell(target_row, 6).number_format = "0.0000"

    emergency = workbook["紧急购电量"]
    if emergency.max_row > 1:
        emergency.delete_rows(2, emergency.max_row - 1)
    target_row = 2
    event_count = 0
    for day in output_dates:
        day_events = merge_emergency_events(physical_by_date.get(day, ()), day, config.parameters.feasibility_tolerance)
        for event_index, event in enumerate(day_events):
            _copy_row_style(emergency, 2, target_row, 3)
            emergency.cell(target_row, 1).value = day if event_index == 0 else None
            emergency.cell(target_row, 2).value = format_event_label(event)
            emergency.cell(target_row, 3).value = round4(event.energy_kwh)
            emergency.cell(target_row, 3).number_format = "0.0000"
            target_row += 1
            event_count += 1

    workbook.save(candidate)
    workbook.close()
    if sha256_file(source) != original_hash:
        raise RuntimeError("official result2.xlsx changed during candidate export")
    return {
        "candidate_path": str(candidate),
        "candidate_sha256": sha256_file(candidate),
        "official_template_path": str(source),
        "official_sha256_before": original_hash,
        "official_sha256_after": sha256_file(source),
        "official_template_unchanged": True,
        "plan_days": len(output_dates),
        "storage_rows": len(output_dates) * 6,
        "emergency_event_rows": event_count,
    }
