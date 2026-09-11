from __future__ import annotations

import csv
import json
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from q1_baseline.run_manifest import environment_snapshot, git_snapshot, sha256_file, timestamp, write_json
from q1_baseline.solver_backend import HighsPyBackend, backend_by_name

from .config import Q2Config, load_config
from .data import Q2InputData, read_q2_inputs
from .evidence import flush_failure_evidence
from .export_validation import validate_result2_candidate
from .exporter import export_result2_candidate
from .forecast import ForecastDay, RecentCompletedSameClockPredictor
from .integrity import verify_input_integrity
from .planner import (
    DailyPlan,
    cold_start_plan,
    compute_b0_plan,
    compute_window_start_soc,
    solve_daily_plan,
)
from .replay import ReplayDay, replay_r0
from .prediction_metrics import build_prediction_observations, evaluate_prediction_observations
from .time_axis import date_range
from .validation import records_to_dict, validate_day


REQUIRED_OUTPUTS = (
    "run_manifest.json", "config_snapshot.json", "model_selection.csv",
    "predictions_load.csv", "predictions_pv.csv", "dispatch_timeseries.csv",
    "daily_metrics.csv", "monthly_metrics.csv", "metrics_summary.csv",
    "prediction_metrics.csv",
    "assertions.json", "solver_days.csv", "solver.log",
    "warnings_and_failures.log", "result2_candidate.xlsx",
    "export_validation.json", "human_feedback.md",
)


def validate_input_only(config_path: str | Path) -> dict[str, Any]:
    config = load_config(Path(config_path))
    integrity = verify_input_integrity(config)
    data = read_q2_inputs(config.normalized_price_input, config.normalized_actual_input, config.parameters)
    from openpyxl import load_workbook

    workbook = load_workbook(config.official_result2_template, read_only=True, data_only=False)
    template_sheet_names = workbook.sheetnames
    plan_sheet = workbook["计划购电量"]
    template_plan_shape = [plan_sheet.max_row, plan_sheet.max_column]
    template_slot1 = plan_sheet.cell(1, 2).value
    template_slot144 = plan_sheet.cell(1, 145).value
    workbook.close()
    if template_sheet_names != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError(f"unexpected official result2 sheets: {template_sheet_names}")
    if template_plan_shape != [335, 147]:
        raise ValueError(f"unexpected official result2 plan shape: {template_plan_shape}")
    if template_slot1 != "0:10-0:20" or template_slot144 != "0:00-0:10+1":
        raise ValueError("official result2 template slot labels changed")
    predictor = RecentCompletedSameClockPredictor(data, config.forecast_version)
    jan1 = predictor.forecast_day(date(2025, 1, 1))
    jan2 = predictor.forecast_day(date(2025, 1, 2))
    jan3 = predictor.forecast_day(date(2025, 1, 3))
    feb1 = predictor.forecast_day(date(2025, 2, 1))
    jan31_slot144 = data.actual_by_template_key[(date(2025, 1, 31), 144)]
    return {
        "passed": True,
        "solver_invoked": False,
        "price_rows": int(data.price.size),
        "actual_rows": len(data.actual),
        "official_result2_sheets": template_sheet_names,
        "official_plan_shape": template_plan_shape,
        "official_slot1": template_slot1,
        "official_slot144": template_slot144,
        "jan1_cold_slots": int(jan1.cold_mask.sum()),
        "jan2_cold_slots": int(jan2.cold_mask.sum()),
        "jan2_slot144_cold": bool(jan2.rows[-1].is_cold_start),
        "jan3_cold_slots": int(jan3.cold_mask.sum()),
        "feb1_all_sources_causal": all(row.source_interval_end is None or row.source_interval_end <= row.decision_time for row in feb1.rows),
        "jan31_slot144_forbidden_at_feb1_decision": jan31_slot144.interval_end > feb1.rows[0].decision_time,
        "input_sha256": {
            "price": integrity["verified_input_sha256"]["price"],
            "actual": integrity["verified_input_sha256"]["actual"],
            "result2_template": integrity["verified_input_sha256"]["official_result2"],
        },
        "input_manifest_sha256": integrity["input_manifest_sha256"],
        "audit_summary_sha256": integrity["audit_summary_sha256"],
        "data_version": integrity["data_version"],
    }


def _write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _period(config: Q2Config, day: date) -> str:
    return "warmup" if day < config.output_start else "formal_output"


def _feedback_template(config: Q2Config, output_dir: Path) -> str:
    return f"""# Q2 人工运行反馈（待填写）

- 当前阶段：5
- 当前状态：Q2待人工正式运行
- 模型版本：{config.model_version}
- 预测版本：{config.forecast_version}
- 数据版本：{config.data_version}
- 运行目录：{output_dir}
- 运行人：待回填
- 运行时间：待回填
- 环境与依赖：待回填
- 365日SOC链及Q2-SOC-BRIDGE-001：待回填
- Q2-COLDSTART-001边界：待回填
- 因果性/leakage断言：待回填
- 三个正式输入SHA-256硬门：待回填
- 预测评价（overall/month/season）：待回填
- Solver逐日状态、MIP gap与日志：待回填
- 全部断言是否通过：待回填
- result2_candidate.xlsx回读是否通过：待回填
- 预热费用与正式费用是否隔离：待回填
- 是否接受为后续人工审查候选：待回填
- WARNING/BLOCKER与人工判断：待回填

注意：运行结果是 HUMAN-RUN / UNREVIEWED CANDIDATE，不是最终比赛结果；
不得覆盖官方 result2.xlsx，不得据此自动进入问题3。
"""


def _prediction_rows(forecasts: list[ForecastDay], kind: str) -> Iterable[dict[str, Any]]:
    for forecast in forecasts:
        for row in forecast.rows:
            yield {
                "template_date": row.template_date.isoformat(),
                "template_slot": row.template_slot,
                "decision_time": row.decision_time.isoformat(sep=" "),
                "information_cutoff": row.information_cutoff.isoformat(sep=" "),
                "interval_start": row.interval_start.isoformat(sep=" "),
                "interval_end": row.interval_end.isoformat(sep=" "),
                "forecast_source": row.forecast_source,
                "forecast_version": row.forecast_version,
                "source_interval_start": "" if row.source_interval_start is None else row.source_interval_start.isoformat(sep=" "),
                "source_interval_end": "" if row.source_interval_end is None else row.source_interval_end.isoformat(sep=" "),
                "predicted_kw": row.load_pred_kw if kind == "load" else row.pv_pred_kw,
            }


PREDICTION_FIELDS = [
    "template_date", "template_slot", "decision_time", "information_cutoff",
    "interval_start", "interval_end", "forecast_source", "forecast_version",
    "source_interval_start", "source_interval_end", "predicted_kw",
]


def _dispatch_rows(
    config: Q2Config,
    data: Q2InputData,
    forecasts: dict[date, ForecastDay],
    replays: list[ReplayDay],
) -> Iterable[dict[str, Any]]:
    for replay in replays:
        period = _period(config, replay.plan.template_date)
        forecast = forecasts[replay.plan.template_date]
        for i, actual in enumerate(replay.actual):
            forecast_row = forecast.rows[i]
            pred_load_kwh = forecast_row.load_pred_kw * config.parameters.delta_t
            pred_pv_kwh = forecast_row.pv_pred_kw * config.parameters.delta_t
            realized_balance = (
                replay.plan.G[i] + actual.pv_kwh + replay.plan.D[i] + replay.emergency_kwh[i]
                - actual.load_kwh - replay.plan.C[i] - replay.surplus_kwh[i]
            )
            soc_residual = (
                replay.plan.S[i + 1] - replay.plan.S[i]
                - config.parameters.charge_efficiency * replay.plan.C[i]
                + replay.plan.D[i] / config.parameters.discharge_efficiency
            )
            yield {
                "period": period,
                "baseline": replay.plan.baseline,
                "template_date": replay.plan.template_date.isoformat(),
                "template_slot": i + 1,
                "clock_slot": actual.natural_clock_minute // 10 + 1,
                "interval_start": actual.interval_start.isoformat(sep=" "),
                "interval_end": actual.interval_end.isoformat(sep=" "),
                "decision_time": forecast_row.decision_time.isoformat(sep=" "),
                "information_cutoff": forecast_row.information_cutoff.isoformat(sep=" "),
                "load_actual_kwh": actual.load_kwh,
                "load_pred_kwh": pred_load_kwh,
                "pv_actual_kwh": actual.pv_kwh,
                "pv_pred_kwh": pred_pv_kwh,
                "price_actual": data.price[i],
                "price_pred": "",
                "grid_plan_kwh": replay.plan.G[i],
                "grid_final_kwh": "",
                "grid_emergency_kwh": replay.emergency_kwh[i],
                "charge_bus_kwh": replay.plan.C[i],
                "discharge_bus_kwh": replay.plan.D[i],
                "spill_kwh": replay.surplus_kwh[i],
                "soc_before_kwh": replay.plan.S[i],
                "soc_after_kwh": replay.plan.S[i + 1],
                "charge_mode": replay.plan.z[i],
                "adjust_down_kwh": "",
                "adjust_up_kwh": "",
                "planned_purchase_cost": replay.planned_purchase_cost[i],
                "downward_adjustment_penalty": "",
                "upward_adjustment_cost": "",
                "emergency_purchase_cost": replay.emergency_purchase_cost[i],
                "total_cost": replay.planned_purchase_cost[i] + replay.emergency_purchase_cost[i],
                "balance_residual": realized_balance,
                "soc_residual": soc_residual,
                "solver_status": replay.plan.status,
            }


DISPATCH_FIELDS = [
    "period", "baseline", "template_date", "template_slot", "clock_slot",
    "interval_start", "interval_end", "decision_time", "information_cutoff",
    "load_actual_kwh", "load_pred_kwh", "pv_actual_kwh", "pv_pred_kwh",
    "price_actual", "price_pred", "grid_plan_kwh", "grid_final_kwh",
    "grid_emergency_kwh", "charge_bus_kwh", "discharge_bus_kwh", "spill_kwh",
    "soc_before_kwh", "soc_after_kwh", "charge_mode", "adjust_down_kwh",
    "adjust_up_kwh", "planned_purchase_cost", "downward_adjustment_penalty",
    "upward_adjustment_cost", "emergency_purchase_cost", "total_cost",
    "balance_residual", "soc_residual", "solver_status",
]


def run_q2(config_path: str | Path, output_dir: str | Path) -> Path:
    """Explicit human-run entry point for the full Jan warmup + Feb-Dec Q2 run."""
    config_path_resolved = Path(config_path).resolve()
    config = load_config(config_path_resolved)
    # Q2-PRE-RUN-PATCH-001 hard gate: no CSV read, prediction, model build,
    # backend initialization, or solver call may occur before this succeeds.
    integrity = verify_input_integrity(config)
    data = read_q2_inputs(config.normalized_price_input, config.normalized_actual_input, config.parameters)
    destination = Path(output_dir).resolve()
    if destination.exists():
        raise FileExistsError(f"refusing to overwrite existing run directory: {destination}")
    destination.mkdir(parents=True)
    write_json(destination / "config_snapshot.json", config.snapshot())
    (destination / "solver.log").write_text("Q2 daily solver evidence\n", encoding="utf-8")
    rounding_warning = (
        "WARNING-Q2-INPUT-ROUNDING-001: normalized CSV kWh values use four decimals; "
        "maximum observed kW/6 difference is about 5.0e-5 kWh. The 5.1e-5 kWh "
        "input-contract tolerance is used only for raw power/energy consistency, "
        "not for MILP feasibility, SOC residual, or balance residual tolerances.\n"
    )
    (destination / "warnings_and_failures.log").write_text(rounding_warning, encoding="utf-8")
    (destination / "human_feedback.md").write_text(_feedback_template(config, destination), encoding="utf-8")
    _write_csv(destination / "model_selection.csv", ["track", "model", "status", "notes"], [
        {"track": "B0", "model": "no_storage_reference", "status": "IMPLEMENTED", "notes": "causal forecast + R0"},
        {"track": "B1", "model": "recent_completed_same_clock + MILP + R0", "status": "IMPLEMENTED", "notes": "locked baseline"},
        {"track": "Candidate", "model": "ForecastProvider interface", "status": "INTERFACE_ONLY", "notes": "no advanced model added"},
    ])
    started_at = timestamp()
    git = git_snapshot(config_path_resolved.parents[2])
    code_hashes = {
        path.name: sha256_file(path)
        for path in sorted(Path(__file__).resolve().parent.glob("*.py"))
    }
    manifest: dict[str, Any] = {
        "run_id": destination.name,
        "question_id": "Q2",
        "stage": 5,
        "status": "HUMAN_RUN_UNREVIEWED",
        "run_status": "HUMAN_RUN_UNREVIEWED",
        "final_competition_result": False,
        "created_at": started_at,
        "started_at": started_at,
        "finished_at": None,
        "model_version": config.model_version,
        "forecast_version": config.forecast_version,
        "data_version": config.data_version,
        "input_paths": {
            "price": str(config.normalized_price_input),
            "actual": str(config.normalized_actual_input),
            "official_result2_template": str(config.official_result2_template),
        },
        "input_sha256": {
            "price": integrity["verified_input_sha256"]["price"],
            "actual": integrity["verified_input_sha256"]["actual"],
            "official_result2_template": integrity["verified_input_sha256"]["official_result2"],
        },
        "input_manifest_sha256": integrity["input_manifest_sha256"],
        "audit_summary_sha256": integrity["audit_summary_sha256"],
        "code_version_or_hashes": code_hashes,
        "config_sha256": sha256_file(config_path_resolved),
        "environment": environment_snapshot(),
        "solver_name": config.solver_name,
        "solver_version": "PENDING_HUMAN_RUN",
        "predictor_versions": {
            "load": config.forecast_version,
            "pv": config.forecast_version,
        },
        "random_seeds": {
            "predictor": "NOT_APPLICABLE_DETERMINISTIC",
            "optimization": "NOT_APPLICABLE_DETERMINISTIC_MILP",
        },
        "information_cutoff_policy": "source_interval_end <= decision_time",
        "output_period": {
            "start": config.output_start.isoformat(),
            "end": config.output_end.isoformat(),
            "template_days": 334,
        },
        "git": git,
        "git_head": git["git_head"],
        "git_dirty": git["git_dirty"],
        "warnings_count": 1,
        "failures_count": 0,
        "warnings": ["WARNING-Q2-INPUT-ROUNDING-001"],
        "last_completed_template_date": None,
        "failed_template_date": None,
        "required_outputs": list(REQUIRED_OUTPUTS),
    }
    write_json(destination / "run_manifest.json", manifest)
    forecasts: list[ForecastDay] = []
    b0_replays: list[ReplayDay] = []
    b1_replays: list[ReplayDay] = []
    assertion_days: list[dict[str, Any]] = []
    daily_metric_rows: list[dict[str, Any]] = []
    solver_rows: list[dict[str, Any]] = []
    active_day: date | None = None
    last_completed_day: date | None = None
    try:
        backend = backend_by_name(config.solver_name)
        if isinstance(backend, HighsPyBackend):
            backend.require_available()
        predictor = RecentCompletedSameClockPredictor(data, config.forecast_version)
        previous_b1: DailyPlan | None = None

        for day in date_range(config.warmup_start, config.output_end):
            active_day = day
            forecast = predictor.forecast_day(day)
            forecasts.append(forecast)
            if previous_b1 is None:
                initial_soc = config.parameters.soc_initial
                bridge_error = 0.0
            else:
                soc_0000 = float(previous_b1.S[-2])
                initial_soc = compute_window_start_soc(
                    soc_0000,
                    float(previous_b1.C[-1]),
                    float(previous_b1.D[-1]),
                    config.parameters,
                )
                bridge_error = abs(initial_soc - float(previous_b1.S[-1]))
                if bridge_error > config.parameters.feasibility_tolerance:
                    raise AssertionError(f"Q2-SOC-BRIDGE-001 failed before {day}")

            actual = tuple(data.actual_by_template_key[(day, slot)] for slot in range(1, 145))
            b0_plan = compute_b0_plan(forecast, data.price, config.parameters, config.parameters.soc_initial)
            b0_replay = replay_r0(b0_plan, actual, data.price, config.parameters)
            b0_replays.append(b0_replay)
            if bool(forecast.cold_mask.all()):
                b1_plan = cold_start_plan(forecast, initial_soc)
            else:
                native_log = destination / f"_solver_native_{day.isoformat()}.log"
                b1_plan = solve_daily_plan(
                    forecast, data.price, config.parameters, initial_soc, backend,
                    native_log,
                )
                with (destination / "solver.log").open("a", encoding="utf-8") as master:
                    master.write(f"\n===== {day.isoformat()} native solver log =====\n")
                    if native_log.is_file():
                        master.write(native_log.read_text(encoding="utf-8", errors="replace"))
                        master.write("\n")
                if native_log.is_file():
                    native_log.unlink()
            manifest["solver_name"] = b1_plan.solver_name
            manifest["solver_version"] = b1_plan.solver_version
            b1_replay = replay_r0(b1_plan, actual, data.price, config.parameters)
            b1_replays.append(b1_replay)
            previous_b1 = b1_plan

            for replay, require_optimal in ((b0_replay, False), (b1_replay, b1_plan.status != "COLD_START_NO_SOLVE")):
                passed, assertions, metrics = validate_day(
                    forecast, replay.plan, replay, data.price, config.parameters,
                    initial_soc if replay.plan.baseline.startswith("B1") else config.parameters.soc_initial,
                    require_optimal,
                )
                assertion_days.append({
                    "template_date": day.isoformat(),
                    "period": _period(config, day),
                    "baseline": replay.plan.baseline,
                    "passed": passed,
                    "soc_bridge_error_kwh": bridge_error if replay.plan.baseline.startswith("B1") else 0.0,
                    "assertions": records_to_dict(assertions),
                })
                daily_metric_rows.append({
                    "period": _period(config, day), "baseline": replay.plan.baseline,
                    "template_date": day.isoformat(), **metrics,
                })
                if not passed:
                    raise RuntimeError(f"daily assertions failed: {day} {replay.plan.baseline}")
            solver_rows.append({
                "template_date": day.isoformat(), "period": _period(config, day),
                "status": b1_plan.status, "solver_name": b1_plan.solver_name,
                "solver_version": b1_plan.solver_version,
                "termination_condition": b1_plan.termination_condition,
                "mip_gap": "" if b1_plan.mip_gap is None else b1_plan.mip_gap,
                "runtime_seconds": b1_plan.runtime_seconds,
            })
            with (destination / "solver.log").open("a", encoding="utf-8") as handle:
                handle.write(f"===== {day.isoformat()} structured evidence =====\n")
                handle.write(json.dumps(solver_rows[-1], ensure_ascii=False) + "\n")
            last_completed_day = day
            manifest["last_completed_template_date"] = day.isoformat()
            _write_csv(
                destination / "solver_days.csv",
                list(solver_rows[0].keys()),
                solver_rows,
            )
            write_json(
                destination / "assertions.json",
                {
                    "passed": all(row["passed"] for row in assertion_days),
                    "incomplete": day != config.output_end,
                    "daily": assertion_days,
                },
            )
            write_json(destination / "run_manifest.json", manifest)

        active_day = None

        _write_csv(destination / "predictions_load.csv", PREDICTION_FIELDS, _prediction_rows(forecasts, "load"))
        _write_csv(destination / "predictions_pv.csv", PREDICTION_FIELDS, _prediction_rows(forecasts, "pv"))
        prediction_observations = build_prediction_observations(
            forecasts, data, config.output_start, config.output_end
        )
        prediction_metric_rows = evaluate_prediction_observations(prediction_observations)
        _write_csv(
            destination / "prediction_metrics.csv",
            list(prediction_metric_rows[0].keys()),
            prediction_metric_rows,
        )
        forecast_by_date = {item.template_date: item for item in forecasts}
        _write_csv(
            destination / "dispatch_timeseries.csv",
            DISPATCH_FIELDS,
            _dispatch_rows(config, data, forecast_by_date, b0_replays + b1_replays),
        )
        metric_fields = list(daily_metric_rows[0].keys())
        _write_csv(destination / "daily_metrics.csv", metric_fields, daily_metric_rows)
        _write_csv(destination / "solver_days.csv", list(solver_rows[0].keys()), solver_rows)

        grouped: dict[tuple[str, str, str], dict[str, float]] = defaultdict(lambda: defaultdict(float))
        for row in daily_metric_rows:
            key = (str(row["period"]), str(row["baseline"]), str(row["template_date"])[:7])
            for name in ("planned_grid_kwh", "charge_kwh", "discharge_kwh", "emergency_kwh", "surplus_kwh", "planned_purchase_cost_yuan", "emergency_purchase_cost_yuan", "total_cost_yuan"):
                grouped[key][name] += float(row[name])
        monthly_rows = [{"period": key[0], "baseline": key[1], "month": key[2], **values} for key, values in sorted(grouped.items())]
        _write_csv(destination / "monthly_metrics.csv", list(monthly_rows[0].keys()), monthly_rows)

        summary_group: dict[tuple[str, str], dict[str, float]] = defaultdict(lambda: defaultdict(float))
        for row in daily_metric_rows:
            key = (str(row["period"]), str(row["baseline"]))
            for name in ("planned_grid_kwh", "charge_kwh", "discharge_kwh", "emergency_kwh", "surplus_kwh", "planned_purchase_cost_yuan", "emergency_purchase_cost_yuan", "total_cost_yuan"):
                summary_group[key][name] += float(row[name])
        summary_wide = [{"period": key[0], "baseline": key[1], **values} for key, values in sorted(summary_group.items())]
        unit_by_metric = {
            "planned_grid_kwh": "kWh", "charge_kwh": "kWh", "discharge_kwh": "kWh",
            "emergency_kwh": "kWh", "surplus_kwh": "kWh",
            "planned_purchase_cost_yuan": "yuan", "emergency_purchase_cost_yuan": "yuan",
            "total_cost_yuan": "yuan",
        }
        summary_rows: list[dict[str, Any]] = []
        for row in summary_wide:
            sample_count = 31 if row["period"] == "warmup" else 334
            baseline = str(row["baseline"])
            for metric, unit in unit_by_metric.items():
                summary_rows.append({
                    "run_id": destination.name,
                    "question_id": "Q2",
                    "model_id": config.model_version if baseline.startswith("B1") else "B0_NO_STORAGE_REFERENCE",
                    "baseline_id": baseline,
                    "split": row["period"],
                    "period": row["period"],
                    "metric": metric,
                    "value": row[metric],
                    "unit": unit,
                    "direction": "min" if metric in {"emergency_kwh", "surplus_kwh", "planned_purchase_cost_yuan", "emergency_purchase_cost_yuan", "total_cost_yuan"} else "descriptive",
                    "sample_count": sample_count,
                    "status": "HUMAN_RUN_UNREVIEWED",
                    "notes": "warmup excluded from formal total" if row["period"] == "warmup" else "formal Q2 output",
                })
        for row in prediction_metric_rows:
            summary_rows.append(
                {
                    "run_id": destination.name,
                    "question_id": "Q2",
                    "model_id": config.forecast_version,
                    "baseline_id": "B1_RECENT_SAME_CLOCK_MILP_R0",
                    "split": "formal_output",
                    "period": f"{row['scope']}:{row['group']}",
                    "metric": f"{row['series']}.{row['metric']}",
                    "value": row["value"],
                    "unit": row["unit"],
                    "direction": "min" if row["metric"] not in {"sample_count", "valid_coverage", "peak_sample_count", "daylight_sample_count", "ramp_sample_count", "peak_threshold_p90"} else "descriptive",
                    "sample_count": row["sample_count"],
                    "status": "POSTHOC_EVALUATION_ONLY",
                    "notes": row["definition"],
                }
            )
        _write_csv(destination / "metrics_summary.csv", list(summary_rows[0].keys()), summary_rows)
        assertion_payload = {
            "passed": all(row["passed"] for row in assertion_days),
            "q2_coldstart_001": {
                "jan1_cold_slots": int(forecasts[0].cold_mask.sum()),
                "jan2_cold_slots": int(forecasts[1].cold_mask.sum()),
                "jan2_slot144_cold": bool(forecasts[1].rows[-1].is_cold_start),
                "jan3_cold_slots": int(forecasts[2].cold_mask.sum()),
            },
            "daily": assertion_days,
        }
        write_json(destination / "assertions.json", assertion_payload)
        b1_tuple = tuple(b1_replays)
        export_record = export_result2_candidate(config, destination / "result2_candidate.xlsx", b1_tuple)
        export_validation = validate_result2_candidate(
            config,
            destination / "result2_candidate.xlsx",
            b1_tuple,
            expected_official_sha256=manifest["input_sha256"]["official_result2_template"],
        )
        write_json(destination / "export_validation.json", export_validation)
        if not export_validation["passed"]:
            raise RuntimeError("result2 candidate reread validation failed")
        formal_b1 = next(row for row in summary_wide if row["period"] == "formal_output" and str(row["baseline"]).startswith("B1"))
        warmup_b1 = next(row for row in summary_wide if row["period"] == "warmup" and str(row["baseline"]).startswith("B1"))
        solved_rows = [row for row in solver_rows if row["solver_name"] != "NONE"]
        finished_at = timestamp()
        manifest.update({
            "run_completed_at": finished_at,
            "finished_at": finished_at,
            "status": "HUMAN_RUN_UNREVIEWED_CANDIDATE",
            "run_status": "HUMAN_RUN_UNREVIEWED_CANDIDATE",
            "solver_name": solved_rows[-1]["solver_name"] if solved_rows else "NONE",
            "solver_version": solved_rows[-1]["solver_version"] if solved_rows else "N/A",
            "assertions_passed": True,
            "export_validation_passed": True,
            "warmup_cost_yuan_separate_not_in_formal": warmup_b1["total_cost_yuan"],
            "formal_q2_total_cost_yuan": formal_b1["total_cost_yuan"],
            "formal_output_days": 334,
            "solver_days": len(solver_rows),
            "failures_count": 0,
            "last_completed_template_date": config.output_end.isoformat(),
            "failed_template_date": None,
            "export": export_record,
            "official_result2_sha256_after": sha256_file(config.official_result2_template),
        })
        write_json(destination / "run_manifest.json", manifest)
        return destination
    except Exception as exc:
        if active_day is not None and not any(
            row.get("template_date") == active_day.isoformat() for row in solver_rows
        ):
            solver_rows.append(
                {
                    "template_date": active_day.isoformat(),
                    "period": _period(config, active_day),
                    "status": "FAILED_OR_INCOMPLETE",
                    "solver_name": config.solver_name,
                    "solver_version": manifest.get("solver_version", "UNKNOWN"),
                    "termination_condition": f"{type(exc).__name__}: {exc}",
                    "mip_gap": "",
                    "runtime_seconds": "",
                }
            )
        if active_day is not None:
            native_log = destination / f"_solver_native_{active_day.isoformat()}.log"
            if native_log.is_file():
                with (destination / "solver.log").open("a", encoding="utf-8") as master:
                    master.write(f"\n===== {active_day.isoformat()} failed native solver log =====\n")
                    master.write(native_log.read_text(encoding="utf-8", errors="replace"))
                    master.write("\n")
        manifest.update({"error_type": type(exc).__name__, "error": str(exc)})
        flush_failure_evidence(
            destination,
            manifest,
            solver_rows,
            assertion_days,
            last_completed_template_date=last_completed_day,
            failed_template_date=active_day,
            error=exc,
        )
        try:
            if forecasts:
                _write_csv(destination / "predictions_load.csv", PREDICTION_FIELDS, _prediction_rows(forecasts, "load"))
                _write_csv(destination / "predictions_pv.csv", PREDICTION_FIELDS, _prediction_rows(forecasts, "pv"))
            if daily_metric_rows:
                _write_csv(destination / "daily_metrics.csv", list(daily_metric_rows[0].keys()), daily_metric_rows)
            available_replays = b0_replays + b1_replays
            if available_replays:
                partial_forecasts = {item.template_date: item for item in forecasts}
                _write_csv(
                    destination / "dispatch_timeseries.csv",
                    DISPATCH_FIELDS,
                    _dispatch_rows(config, data, partial_forecasts, available_replays),
                )
        except Exception as evidence_exc:
            with (destination / "warnings_and_failures.log").open("a", encoding="utf-8") as handle:
                handle.write(
                    f"{timestamp()} SECONDARY_EVIDENCE_FLUSH_FAILURE "
                    f"{type(evidence_exc).__name__}: {evidence_exc}\n"
                )
        raise
