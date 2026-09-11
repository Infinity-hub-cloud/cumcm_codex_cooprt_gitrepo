from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from q1_baseline.run_manifest import sha256_file

from .config import Q2Config
from .exporter import (
    FOUR_HOUR_LABELS,
    format_event_label,
    index_physical_by_natural_date,
    merge_emergency_events,
    natural_day_four_hour_summary,
    physical_dispatch_rows,
    round4,
)
from .replay import ReplayDay
from .time_axis import date_range


def validate_result2_candidate(
    config: Q2Config,
    candidate_path: Path,
    b1_replays: tuple[ReplayDay, ...],
    expected_official_sha256: str | None = None,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    official = config.official_result2_template.resolve()
    candidate = candidate_path.resolve()
    checks: list[dict[str, Any]] = []

    def add(name: str, passed: bool, detail: Any) -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    add("candidate_is_separate_file", candidate != official, str(candidate))
    add("candidate_exists", candidate.is_file(), str(candidate))
    if not candidate.is_file():
        return {"passed": False, "checks": checks}
    replay_by_date = {item.plan.template_date: item for item in b1_replays}
    output_dates = date_range(config.output_start, config.output_end)
    physical = physical_dispatch_rows(b1_replays)
    physical_by_date = index_physical_by_natural_date(physical)
    # Normal mode is deliberate: thousands of random cell reads on a read-only
    # worksheet repeatedly rescan its XML stream and are prohibitively slow.
    workbook = load_workbook(candidate, read_only=False, data_only=False)
    add("sheet_names", workbook.sheetnames == ["计划购电量", "充放电量", "紧急购电量"], workbook.sheetnames)
    plan = workbook["计划购电量"]
    storage = workbook["充放电量"]
    emergency = workbook["紧急购电量"]
    add("plan_dimensions", (plan.max_row, plan.max_column) == (335, 147), [plan.max_row, plan.max_column])
    add("storage_dimensions", (storage.max_row, storage.max_column) == (2005, 6), [storage.max_row, storage.max_column])
    add("slot_labels_not_rotated", plan.cell(1, 2).value == "0:10-0:20" and plan.cell(1, 145).value == "0:00-0:10+1", [plan.cell(1, 2).value, plan.cell(1, 145).value])

    def as_date(value: Any) -> date | None:
        if isinstance(value, datetime):
            return value.date()
        return value if isinstance(value, date) else None

    official_workbook = load_workbook(official, read_only=True, data_only=False)
    official_plan = official_workbook["计划购电量"]
    candidate_headers = [plan.cell(1, column).value for column in range(1, 148)]
    official_headers = [official_plan.cell(1, column).value for column in range(1, 148)]
    add("all_official_plan_headers_preserved", candidate_headers == official_headers, candidate_headers)
    official_workbook.close()

    plan_error = 0.0
    blank_count = 0
    plan_date_errors = 0
    for row_index, day in enumerate(output_dates, start=2):
        replay = replay_by_date[day]
        if as_date(plan.cell(row_index, 1).value) != day:
            plan_date_errors += 1
        for slot in range(144):
            value = plan.cell(row_index, slot + 2).value
            if not isinstance(value, (int, float)):
                blank_count += 1
            else:
                plan_error = max(plan_error, abs(float(value) - round4(replay.plan.G[slot])))
        plan_error = max(plan_error, abs(float(plan.cell(row_index, 146).value) - round4(sum(replay.plan.G))))
        plan_error = max(plan_error, abs(float(plan.cell(row_index, 147).value) - round4(sum(replay.planned_purchase_cost))))
    add("plan_no_blanks", blank_count == 0, blank_count)
    add("plan_dates", plan_date_errors == 0, plan_date_errors)
    add("plan_values_original_slot_order", plan_error <= 5e-8, plan_error)

    storage_error = 0.0
    storage_labels_ok = True
    storage_dates_ok = True
    for date_index, day in enumerate(output_dates):
        summaries = natural_day_four_hour_summary(physical_by_date.get(day, ()), day)
        for group, label in enumerate(FOUR_HOUR_LABELS):
            row = 2 + date_index * 6 + group
            storage_labels_ok &= storage.cell(row, 2).value == label
            storage_dates_ok &= (
                as_date(storage.cell(row, 1).value) == day if group == 0
                else storage.cell(row, 1).value is None
            )
            storage_error = max(
                storage_error,
                abs(float(storage.cell(row, 3).value) - round4(summaries[group][0])),
                abs(float(storage.cell(row, 4).value) - round4(summaries[group][1])),
            )
        soc0 = round4(replay_by_date[day - timedelta(days=1)].plan.S[-2])
        soc24 = round4(replay_by_date[day].plan.S[-2])
        storage_error = max(
            storage_error,
            abs(float(storage.cell(2 + date_index * 6, 6).value) - soc0),
            abs(float(storage.cell(3 + date_index * 6, 6).value) - soc24),
        )
    add("natural_day_four_hour_labels", storage_labels_ok, storage_labels_ok)
    add("natural_day_storage_dates", storage_dates_ok, storage_dates_ok)
    add("natural_day_storage_and_soc", storage_error <= 5e-8, storage_error)

    expected_events_with_group_index = [
        (event, event_index)
        for day in output_dates
        for event_index, event in enumerate(
            merge_emergency_events(physical_by_date.get(day, ()), day, config.parameters.feasibility_tolerance)
        )
    ]
    expected_events = [item[0] for item in expected_events_with_group_index]
    emergency_error = 0.0
    actual_event_rows = max(0, emergency.max_row - 1)
    if actual_event_rows == len(expected_events):
        for i, event in enumerate(expected_events, start=2):
            emergency_error = max(emergency_error, abs(float(emergency.cell(i, 3).value) - round4(event.energy_kwh)))
    else:
        emergency_error = float("inf")
    add("emergency_event_count", actual_event_rows == len(expected_events), [actual_event_rows, len(expected_events)])
    add("emergency_event_energy", emergency_error <= 5e-8, emergency_error)
    event_metadata_ok = actual_event_rows == len(expected_events) and all(
        (as_date(emergency.cell(i, 1).value) == event.natural_date if event_index == 0 else emergency.cell(i, 1).value is None)
        and emergency.cell(i, 2).value == format_event_label(event)
        for i, (event, event_index) in enumerate(expected_events_with_group_index, start=2)
    )
    add("emergency_event_dates_and_labels", event_metadata_ok, event_metadata_ok)
    physical_emergency_total = sum(
        row.emergency_kwh
        for day in output_dates
        for row in physical_by_date.get(day, ())
        if row.emergency_kwh > config.parameters.feasibility_tolerance
    )
    event_total_error = abs(sum(event.energy_kwh for event in expected_events) - physical_emergency_total)
    add("emergency_merge_preserves_total", event_total_error <= config.parameters.aggregate_tolerance, event_total_error)

    formula_cells: list[str] = []
    ellipsis_cells: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{sheet.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and "⁝" in cell.value:
                    ellipsis_cells.append(f"{sheet.title}!{cell.coordinate}")
    add("no_formulas", not formula_cells, formula_cells)
    add("no_template_ellipsis", not ellipsis_cells, ellipsis_cells)
    workbook.close()
    official_hash = sha256_file(official)
    if expected_official_sha256 is None:
        add("official_template_present", official.is_file(), str(official))
    else:
        add("official_template_unchanged", official_hash == expected_official_sha256, {"before": expected_official_sha256, "after": official_hash})
    return {
        "passed": all(item["passed"] for item in checks),
        "checks": checks,
        "candidate_sha256": sha256_file(candidate),
        "official_sha256": official_hash,
    }
