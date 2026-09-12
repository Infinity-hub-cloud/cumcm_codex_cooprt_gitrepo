"""Read-only source audit; writes diagnostic evidence only, never solves."""
import csv
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
RAW = BASE.parent / 'C题' / '附件'
OUT = BASE / '审计' / 'Q3_独立审计_20260912'

def write(name, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

def period(t):
    h = t.hour + t.minute / 60
    return 'morning_06_10' if 6 <= h < 10 else 'noon_10_14' if 10 <= h < 14 else 'afternoon_14_18' if 14 <= h < 18 else 'night_18_06'

def main():
    OUT.mkdir(parents=True, exist_ok=False)
    fingerprints = []
    for p in [RAW / '附件2.xlsx', RAW / '附件3.xlsx', RAW.parent / 'C题.pdf']:
        fingerprints.append({'path': str(p), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()})
    w = load_workbook(RAW / '附件2.xlsx', read_only=True, data_only=True)
    values = list(w.worksheets[1].values)
    actual = {}
    for row in values[1:]:
        d = row[0]
        for i, v in enumerate(row[1:], 1):
            actual[d + timedelta(minutes=10*i)] = float(v)
    w.close()
    w = load_workbook(RAW / '附件3.xlsx', read_only=True, data_only=True)
    raw = list(w.worksheets[0].values)
    groups = defaultdict(list)
    samples = []
    raw_points = {}
    for row in raw[1:]:
        if row[0]:
            day = datetime.strptime(str(row[0]), '%Y-%m-%d')
        hour = int(str(row[1]).split(':')[0])
        issue = day + timedelta(hours=hour)
        for lead, v in enumerate(row[2:], 1):
            raw_points[(issue, lead)] = float(v)
            t = issue + timedelta(hours=lead)
            old = t - timedelta(hours=1)
            av = [actual.get(old + timedelta(minutes=10*k)) for k in range(6)]
            if t not in actual or old not in actual or None in av:
                continue
            vals = {'H0_hour_start':actual[old], 'H0_hour_mean':float(np.mean(av)), 'H1_point':actual[t]}
            for hypothesis, truth in vals.items():
                for category in ['ALL', period(t)]:
                    for month in ['ALL', f'{t.month:02}']:
                        groups[(hypothesis, hour, category, month)].append((float(v), truth))
            if day.strftime('%m-%d') in ['01-01','03-20','06-21','09-23','12-21']:
                samples.append({'issue':issue.isoformat(), 'lead':lead, 'target_H1':t.isoformat(), 'forecast_kw':v, **vals})
    w.close()
    metrics = []
    for (hyp,hour,cat,month), pairs in sorted(groups.items()):
        a = np.array(pairs); e = a[:,0]-a[:,1]
        metrics.append({'hypothesis':hyp,'issue_hour':hour,'period':cat,'month':month,'n':len(e),'MAE_kw':float(np.mean(abs(e))),'RMSE_kw':float(np.sqrt(np.mean(e*e))),'bias_kw':float(np.mean(e))})
    write('point_semantics_metrics.csv', metrics)
    write('raw_examples.csv', samples)
    # Same target set: both mappings supported by adjacent points of same issue.
    mapping_groups = defaultdict(list)
    for (issue, lead), left in raw_points.items():
        if lead == 24:
            continue
        right = raw_points[(issue,lead+1)]
        for k in range(6):
            t = issue+timedelta(hours=lead,minutes=10*k)
            if t not in actual:
                continue
            for method,pred in [('H1_ZOH',left),('H1_INTERP',left+(right-left)*k/6)]:
                for cat in ['ALL',period(t)]:
                    mapping_groups[(method,issue.hour,cat)].append((pred,actual[t]))
    mapping_metrics=[]
    for (method,hour,cat),pairs in sorted(mapping_groups.items()):
        a=np.array(pairs);e=a[:,0]-a[:,1]
        mapping_metrics.append({'method':method,'issue_hour':hour,'period':cat,'n':len(e),'MAE_kw':float(np.mean(abs(e))),'RMSE_kw':float(np.sqrt(np.mean(e*e))),'bias_kw':float(np.mean(e))})
    write('ten_minute_common_support_metrics.csv',mapping_metrics)
    normalized=BASE/'预处理审计输出_时间口径修订版'/'normalized'
    mismatches=0
    with (normalized/'attachment3_forecast_long.csv').open(encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            issue=datetime.fromisoformat(r['issue_date'])+timedelta(hours=int(r['issue_time'].split(':')[0]))
            mismatches+=float(r['forecast_kw'])!=raw_points[(issue,int(r['lead_hour']))]
    with (normalized/'attachment2_actual_long.csv').open(encoding='utf-8-sig') as f:
        reader=csv.DictReader(f)
        rows=list(reader)
    actual_fields=list(rows[0])
    summary={'raw_a3_header':list(raw[0]),'raw_a3_rows':len(raw)-1,'raw_a2_pv_points':len(actual),'normalized_a3_value_mismatches':mismatches,'actual_csv_fields':actual_fields,'source_hashes':fingerprints,'purpose':'diagnostic hindsight comparison only; no solver; no forecast policy tuning','comparison_support':'H0/H1 same issue-lead records; ZOH/INTERP same physical targets, lead1..23; excludes unsupported first hour and beyond lead24'}
    (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=True))
    for r in metrics:
        if r['period']=='ALL' and r['month']=='ALL':print(r)
    for r in mapping_metrics:
        if r['period']=='ALL': print(r)

if __name__=='__main__':
    main()
