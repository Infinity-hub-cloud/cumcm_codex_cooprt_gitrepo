from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from q1_baseline.run_manifest import sha256_file
from q2_baseline.exporter import (
    FOUR_HOUR_LABELS, PhysicalDispatch, format_event_label, index_physical_by_natural_date,
    merge_emergency_events, natural_day_four_hour_summary, round4,
)
from q2_baseline.time_axis import date_range

from .config import Q3Config
from .planner import RollingPlan
from .state_machine import ExecutedInterval


def validate_result3_candidate(
    config: Q3Config,
    candidate_path: Path,
    plans: dict[date, RollingPlan],
    executed: tuple[ExecutedInterval, ...],
    expected_official_sha256: str,
) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []

    def add(name: str, passed: bool, detail: Any) -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    official = config.official_result3_template.resolve()
    candidate = candidate_path.resolve()
    add("candidate_is_separate_file", candidate != official, str(candidate))
    add("candidate_exists", candidate.is_file(), str(candidate))
    if not candidate.is_file():
        return {"passed": False, "checks": checks}
    workbook = load_workbook(candidate, read_only=False, data_only=False)
    official_workbook = load_workbook(official, read_only=True, data_only=False)
    expected_sheets = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    add("sheet_names", workbook.sheetnames == expected_sheets, workbook.sheetnames)
    add("all_144_slot_labels", [workbook["计划购电量"].cell(1, c).value for c in range(2, 146)] == [official_workbook["计划购电量"].cell(1, c).value for c in range(2, 146)], "columns 2..145")
    output_days = date_range(config.output_start, config.output_end)
    by_template = {(date.fromisoformat(row.template_date), row.template_slot): row for row in executed}

    def as_date(value: Any) -> date | None:
        if isinstance(value, datetime):
            return value.date()
        return value if isinstance(value, date) else None

    for sheet_name, field in (("计划购电量", "G"), ("调整购电量", "Q")):
        sheet = workbook[sheet_name]
        add(f"{sheet_name}_dimensions", (sheet.max_row, sheet.max_column) == (335, 147), [sheet.max_row, sheet.max_column])
        add(
            f"{sheet_name}_slot_order",
            sheet.cell(1, 2).value == "0:10-0:20" and sheet.cell(1, 145).value == "0:00-0:10+1",
            [sheet.cell(1, 2).value, sheet.cell(1, 145).value],
        )
        errors = 0
        max_error = 0.0
        for row_index, day in enumerate(output_days, start=2):
            if as_date(sheet.cell(row_index, 1).value) != day:
                errors += 1
            vector = plans[day].G if field == "G" else plans[day].Q
            for i, expected in enumerate(vector, start=2):
                value = sheet.cell(row_index, i).value
                if not isinstance(value, (int, float)):
                    errors += 1
                else:
                    max_error = max(max_error, abs(float(value) - round4(float(expected))))
            rows = [by_template[(day, slot)] for slot in range(1, 145)]
            expected_quantity = round4(sum(float(value) for value in vector))
            expected_cost = round4(
                sum(
                    row.planned_purchase_cost if field == "G" else row.regular_purchase_cost
                    for row in rows
                )
            )
            max_error = max(
                max_error,
                abs(float(sheet.cell(row_index, 146).value) - expected_quantity),
                abs(float(sheet.cell(row_index, 147).value) - expected_cost),
            )
        add(f"{sheet_name}_all_values", errors == 0 and max_error <= 5e-8, {"errors": errors, "max_error": max_error})

    physical = tuple(PhysicalDispatch(row.interval_start, row.interval_end, row.C, row.D, row.E) for row in sorted(executed, key=lambda item: item.interval_start))
    physical_by_day = index_physical_by_natural_date(physical)
    executed_by_start = {row.interval_start: row for row in executed}
    storage = workbook["充放电量"]
    add("storage_dimensions", (storage.max_row, storage.max_column) == (1 + len(output_days) * 6, 6), [storage.max_row, storage.max_column])
    storage_errors = 0
    storage_max_error = 0.0
    for date_index, day in enumerate(output_days):
        summaries = natural_day_four_hour_summary(physical_by_day.get(day, ()), day)
        first = executed_by_start[datetime.combine(day, datetime.min.time())]
        last = executed_by_start[datetime.combine(day, datetime.min.time()) + timedelta(hours=23, minutes=50)]
        for group, label in enumerate(FOUR_HOUR_LABELS):
            row_index = 2 + date_index * 6 + group
            expected_date = day if group == 0 else None
            expected_time = datetime.min.time() if group == 0 else "24:00" if group == 1 else None
            expected_soc = round4(first.soc_before) if group == 0 else round4(last.soc_after) if group == 1 else None
            if as_date(storage.cell(row_index, 1).value) != expected_date or storage.cell(row_index, 2).value != label or storage.cell(row_index, 5).value != expected_time:
                storage_errors += 1
            for column, expected in ((3, summaries[group][0]), (4, summaries[group][1])):
                value = storage.cell(row_index, column).value
                if not isinstance(value, (int, float)):
                    storage_errors += 1
                else:
                    storage_max_error = max(storage_max_error, abs(float(value) - round4(expected)))
            soc_value = storage.cell(row_index, 6).value
            if expected_soc is None:
                storage_errors += soc_value is not None
            elif not isinstance(soc_value, (int, float)):
                storage_errors += 1
            else:
                storage_max_error = max(storage_max_error, abs(float(soc_value) - expected_soc))
    add("storage_all_values", storage_errors == 0 and storage_max_error <= 5e-8, {"errors": storage_errors, "max_error": storage_max_error})

    expected_events = []
    output_day_set = set(output_days)
    for day in output_days:
        events = merge_emergency_events(physical_by_day.get(day, ()), day, config.parameters.feasibility_tolerance)
        for index, event in enumerate(events):
            expected_events.append((day if index == 0 else None, format_event_label(event), round4(event.energy_kwh)))
    emergency = workbook["紧急购电量"]
    actual_events = []
    for row_index in range(2, emergency.max_row + 1):
        actual_events.append((as_date(emergency.cell(row_index, 1).value), emergency.cell(row_index, 2).value, emergency.cell(row_index, 3).value))
    add("emergency_dimensions", emergency.max_column == 3 and emergency.max_row == 1 + len(expected_events), [emergency.max_row, emergency.max_column])
    event_errors = abs(len(actual_events) - len(expected_events))
    event_max_error = 0.0
    for actual_row, expected_row in zip(actual_events, expected_events):
        if actual_row[:2] != expected_row[:2] or not isinstance(actual_row[2], (int, float)):
            event_errors += 1
        else:
            event_max_error = max(event_max_error, abs(float(actual_row[2]) - expected_row[2]))
    merged = sum(event[2] for event in expected_events)
    raw = sum(row.E for row in executed if row.interval_start.date() in output_day_set)
    add("emergency_all_values_and_merge_conservation", event_errors == 0 and event_max_error <= 5e-8 and abs(merged - raw) <= len(expected_events) * 5.1e-5, {"errors": event_errors, "max_error": event_max_error, "raw_kwh": raw, "merged_rounded_kwh": merged})

    formula_cells: list[str] = []
    ellipsis_cells: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{sheet.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and "⁝" in cell.value:
                    ellipsis_cells.append(f"{sheet.title}!{cell.coordinate}")
    add("no_formulas", not formula_cells, formula_cells)
    add("no_template_ellipsis", not ellipsis_cells, ellipsis_cells)
    workbook.close()
    official_workbook.close()
    official_hash = sha256_file(official)
    add("official_template_unchanged", official_hash == expected_official_sha256, official_hash)
    return {
        "passed": all(item["passed"] for item in checks),
        "checks": checks,
        "candidate_sha256": sha256_file(candidate),
        "official_sha256": official_hash,
    }
