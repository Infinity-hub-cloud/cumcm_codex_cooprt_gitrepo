from __future__ import annotations

import shutil
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from q1_baseline.run_manifest import sha256_file
from q2_baseline.exporter import (
    FOUR_HOUR_LABELS,
    PhysicalDispatch,
    format_event_label,
    index_physical_by_natural_date,
    merge_emergency_events,
    natural_day_four_hour_summary,
    round4,
)
from q2_baseline.time_axis import date_range

from .config import Q3Config
from .planner import RollingPlan
from .state_machine import ExecutedInterval


def _copy_row_style(sheet: Any, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)


def _physical(executed: Iterable[ExecutedInterval]) -> tuple[PhysicalDispatch, ...]:
    return tuple(
        PhysicalDispatch(row.interval_start, row.interval_end, row.C, row.D, row.E)
        for row in sorted(executed, key=lambda item: item.interval_start)
    )


def export_result3_candidate(
    config: Q3Config,
    candidate_path: Path,
    plans: dict[date, RollingPlan],
    executed: tuple[ExecutedInterval, ...],
) -> dict[str, Any]:
    source = config.official_result3_template.resolve()
    candidate = candidate_path.resolve()
    if source == candidate:
        raise ValueError("candidate path must differ from official result3.xlsx")
    if candidate.exists():
        raise FileExistsError(f"refusing to overwrite candidate: {candidate}")
    original_hash = sha256_file(source)
    expected_days = date_range(config.warmup_start, config.output_end)
    if tuple(sorted(plans)) != expected_days:
        raise ValueError("result3 export requires the complete Jan1-Dec31 plan chain")

    candidate.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, candidate)
    workbook = load_workbook(candidate, data_only=False)
    expected_sheets = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(f"unexpected result3 sheets: {workbook.sheetnames}")

    output_days = date_range(config.output_start, config.output_end)
    executed_by_template = {(date.fromisoformat(row.template_date), row.template_slot): row for row in executed}
    for sheet_name, field in (("计划购电量", "G"), ("调整购电量", "Q")):
        sheet = workbook[sheet_name]
        if (sheet.max_row, sheet.max_column) != (335, 147):
            raise ValueError(f"official {sheet_name} dimensions changed")
        if sheet.cell(1, 2).value != "0:10-0:20" or sheet.cell(1, 145).value != "0:00-0:10+1":
            raise ValueError(f"official {sheet_name} slot labels changed")
        for row_index, day in enumerate(output_days, start=2):
            plan = plans[day]
            values = plan.G if field == "G" else plan.Q
            sheet.cell(row_index, 1).value = day
            for i, value in enumerate(values, start=2):
                sheet.cell(row_index, i).value = round4(float(value))
                sheet.cell(row_index, i).number_format = "0.0000"
            rows = [executed_by_template[(day, slot)] for slot in range(1, 145)]
            sheet.cell(row_index, 146).value = round4(sum(float(value) for value in values))
            if field == "G":
                daily_cost = sum(row.planned_purchase_cost for row in rows)
            else:
                daily_cost = sum(row.regular_purchase_cost for row in rows)
            sheet.cell(row_index, 147).value = round4(daily_cost)
            sheet.cell(row_index, 146).number_format = "0.0000"
            sheet.cell(row_index, 147).number_format = "0.0000"

    physical_by_day = index_physical_by_natural_date(_physical(executed))
    storage = workbook["充放电量"]
    if storage.max_row > 1:
        storage.delete_rows(2, storage.max_row - 1)
    executed_by_start = {row.interval_start: row for row in executed}
    for date_index, day in enumerate(output_days):
        natural = physical_by_day.get(day, ())
        summaries = natural_day_four_hour_summary(natural, day)
        first = executed_by_start[datetime.combine(day, time.min)]
        last = executed_by_start[
            datetime.combine(day, time.min) + timedelta(hours=23, minutes=50)
        ]
        for group, label in enumerate(FOUR_HOUR_LABELS):
            row = 2 + date_index * 6 + group
            _copy_row_style(storage, 2, row, 6)
            storage.cell(row, 1).value = day if group == 0 else None
            storage.cell(row, 2).value = label
            storage.cell(row, 3).value = round4(summaries[group][0])
            storage.cell(row, 4).value = round4(summaries[group][1])
            storage.cell(row, 3).number_format = "0.0000"
            storage.cell(row, 4).number_format = "0.0000"
            if group == 0:
                storage.cell(row, 5).value = time.min
                storage.cell(row, 6).value = round4(first.soc_before)
            elif group == 1:
                storage.cell(row, 5).value = "24:00"
                storage.cell(row, 6).value = round4(last.soc_after)

    emergency = workbook["紧急购电量"]
    if emergency.max_row > 1:
        emergency.delete_rows(2, emergency.max_row - 1)
    row_index = 2
    event_count = 0
    for day in output_days:
        events = merge_emergency_events(
            physical_by_day.get(day, ()), day, config.parameters.feasibility_tolerance
        )
        for event_index, event in enumerate(events):
            _copy_row_style(emergency, 2, row_index, 3)
            emergency.cell(row_index, 1).value = day if event_index == 0 else None
            emergency.cell(row_index, 2).value = format_event_label(event)
            emergency.cell(row_index, 3).value = round4(event.energy_kwh)
            emergency.cell(row_index, 3).number_format = "0.0000"
            row_index += 1
            event_count += 1

    workbook.save(candidate)
    workbook.close()
    if sha256_file(source) != original_hash:
        raise RuntimeError("official result3.xlsx changed during candidate export")
    return {
        "candidate_path": str(candidate),
        "candidate_sha256": sha256_file(candidate),
        "official_template_path": str(source),
        "official_sha256_before": original_hash,
        "official_sha256_after": sha256_file(source),
        "official_template_unchanged": True,
        "plan_days": len(output_days),
        "storage_rows": len(output_days) * 6,
        "emergency_event_rows": event_count,
    }
