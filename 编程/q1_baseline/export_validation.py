from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np

from .baseline import DispatchSolution
from .export_templates import natural_day_four_hour_summaries, round4, sha256_file
from .parameters import Q1Parameters
from .time_index import build_official_q1_intervals


def validate_result1_candidate(
    official_template: str | Path,
    candidate_path: str | Path,
    solution: DispatchSolution,
    params: Q1Parameters,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    official = Path(official_template).resolve()
    candidate = Path(candidate_path).resolve()
    checks: list[dict[str, Any]] = []

    def add(name: str, passed: bool, detail: Any) -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    add("candidate_is_separate_file", candidate != official, str(candidate))
    add("candidate_exists", candidate.is_file(), str(candidate))
    if not candidate.is_file():
        return {"passed": False, "checks": checks}

    workbook = load_workbook(candidate, read_only=True, data_only=False)
    add("sheet_names", workbook.sheetnames == ["计划购电量", "充放电量"], workbook.sheetnames)
    plan = workbook["计划购电量"]
    storage = workbook["充放电量"]
    add("plan_dimensions", (plan.max_row, plan.max_column) == (145, 2), [plan.max_row, plan.max_column])
    add("storage_dimensions", (storage.max_row, storage.max_column) == (7, 5), [storage.max_row, storage.max_column])

    expected_intervals = build_official_q1_intervals(params.interval_count, 10)
    labels = [plan.cell(row, 1).value for row in range(2, 146)]
    expected_labels = [item.official_template_label for item in expected_intervals]
    add("all_144_labels", labels == expected_labels, {"first": labels[:2], "last": labels[-2:]})
    plan_values = [plan.cell(row, 2).value for row in range(2, 146)]
    add("plan_no_blanks", all(isinstance(v, (int, float)) for v in plan_values), sum(v is None for v in plan_values))
    if all(isinstance(v, (int, float)) for v in plan_values):
        expected_values = np.asarray([round4(value) for value in solution.G])
        error = float(np.max(np.abs(np.asarray(plan_values, dtype=float) - expected_values)))
        add("slot_value_mapping", error <= 5e-8, error)
    formats = [plan.cell(row, 2).number_format for row in range(2, 146)]
    add("plan_four_decimal_format", all(fmt == "0.0000" for fmt in formats), sorted(set(formats)))

    charge = natural_day_four_hour_summaries(solution.C)
    discharge = natural_day_four_hour_summaries(solution.D)
    summary_error = 0.0
    for row in range(2, 8):
        summary_error = max(
            summary_error,
            abs(float(storage.cell(row, 2).value) - round4(charge[row - 2])),
            abs(float(storage.cell(row, 3).value) - round4(discharge[row - 2])),
        )
    add("natural_day_storage_summaries", summary_error <= 5e-8, summary_error)
    midnight = round4(solution.S[params.interval_count - 1])
    add("soc_0000_2400", storage["E2"].value == midnight and storage["E3"].value == midnight, [storage["E2"].value, storage["E3"].value])

    formula_cells = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{sheet.title}!{cell.coordinate}")
    add("no_formulas", not formula_cells, formula_cells)
    workbook.close()
    add("official_template_present", official.is_file(), str(official))
    return {
        "passed": all(check["passed"] for check in checks),
        "checks": checks,
        "candidate_sha256": sha256_file(candidate),
        "official_sha256": sha256_file(official),
    }

