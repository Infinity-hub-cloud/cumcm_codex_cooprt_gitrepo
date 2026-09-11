from __future__ import annotations

from datetime import time
from pathlib import Path
from typing import Any


def _sample_minute(value: Any) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    if text == "0:00+1":
        return 1440
    hour_text, minute_text = text.split(":", 1)
    return int(hour_text) * 60 + int(minute_text)


def read_raw_attachment1_authority_check(path: str | Path) -> list[dict[str, Any]]:
    """Read the official workbook without modifying it.

    The optimization runner intentionally consumes the audited normalized CSV;
    this function exists only for an explicit human authority cross-check.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("openpyxl is required for raw authority checks") from exc

    source = Path(path).resolve()
    workbook = load_workbook(source, read_only=True, data_only=False)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"unexpected attachment1 sheets: {workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        raw_time, price, load_kw, pv_kw = values
        rows.append(
            {
                "excel_row": excel_row,
                "sample_minute": _sample_minute(raw_time),
                "price": float(price),
                "load_kw": float(load_kw),
                "pv_kw": float(pv_kw),
            }
        )
    workbook.close()
    if len(rows) != 144 or [row["sample_minute"] for row in rows] != list(
        range(10, 1450, 10)
    ):
        raise ValueError("official attachment1 time axis failed authority check")
    return rows

